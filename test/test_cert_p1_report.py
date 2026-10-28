# -*- coding: utf-8 -*-
"""
認證測試報告 P1 —— Excel 報告格式離線驗證
==========================================
以合成資料產生報告到暫存資料夾（測試結束即刪除），檢查：
頁籤順序與動態數量、欄位、Excel Table 樣式、圖表（系列、缺值斷線、plotVisOnly、h:mm）、
Summary 蛇形圖表＋真正箭頭（休息時間）、圖表下方不放表格與說明、格線淡灰、待截圖標示（不含任何圖片）、N/A 與有效 0、
Raw Data 可追溯、操作人員確認不改寫原始資料、重複產生結果一致。
不連設備、不讀寫單次報告的任何檔案。
執行：python test_cert_p1_report.py
"""

import os
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import cert_report_config as CFG
from cert_stage_engine import Obs, analyze
from cert_report_xlsx import (LOG_HEADERS, _cn_num, arrow_text, overall_title, rest_minutes_text, save_report,
                              summary_titles)

_results = []


def check(label, ok):
    _results.append(bool(ok))
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}")


T0 = datetime(2026, 1, 5, 10, 0, 0)
CUR = {"D": -140.0, "C": 135.0, "I": 0.0}


def obs_at(sec, k, **v):
    t = T0 + timedelta(seconds=sec)
    if k == "X":
        return Obs(t=t, ok=False, fail_reason="模擬讀取失敗", src=f"syn:{sec}")
    vals = {"soc": 50, "current_a": CUR[k], "rack_max_c": 40 + sec / 600, "rack_min_c": 30,
            "voltage_v": 1200.0, "sensor_c": 25.0, "pcs_module_c": None, "power_kw": None}
    if k == "D":                                      # 放電段有 Cell 電壓（原始 4 位小數，確認不被降精度）；
        vals.update(cell_vmax_v=3.4015, cell_vmin_v=3.3852)   # 其餘（充電／休息）模擬舊資料沒有此欄
    vals.update(v)
    return Obs(t=t, values=vals, src=f"syn:{sec}")


def scenario():
    o = []
    for s in range(0, 1201, 10):                       # 10:00～10:20 放電，10:10:00～10:10:40 停頓
        if 600 <= s <= 640:
            o.append(obs_at(s, "I"))
        elif 290 <= s <= 310:
            o.append(obs_at(s, "D", sensor_c=None))     # 10:05 前後傳感器無值 → N/A
        elif 890 <= s <= 910:
            o.append(obs_at(s, "D", soc=0))             # 10:15 前後 SOC 有效 0
        else:
            o.append(obs_at(s, "D"))
    o += [obs_at(s, "I") for s in range(1210, 1501, 10)]                 # 休息
    for s in range(1510, 2401, 10):                                       # 10:25:10～10:40 充電
        o.append(obs_at(s, "X") if 1780 <= s <= 1820 else obs_at(s, "C"))  # 10:29:40～10:30:20 讀取失敗
    o += [obs_at(s, "I") for s in range(2410, 2521, 10)]
    o += [obs_at(s, "D") for s in range(2530, 2571, 10)]                 # 10:42:10～10:42:50 短放電
    o += [obs_at(s, "I") for s in range(2580, 2701, 10)]
    return o


META = {"conditions": ["冷氣溫度：測試", "硬體概述：測試", "測試條件：測試"],
        "source_notes": ["合成資料"], "generated_at": "2026-01-05 12:00:00",
        "cabinet_spec": "250KW", "cycles": 2}         # 刻意不用 340KW：確認規格來自測試資訊、沒有寫死

tmp = tempfile.mkdtemp(prefix="cert_p1_")
try:
    p = CFG.params()
    res = analyze(scenario(), p)
    path = save_report(res, META, os.path.join(tmp, "r.xlsx"))
    wb = load_workbook(path)

    print("\n== 頁籤與階段 ==")
    check("階段判定：第1次放電（含 1 次確定停頓）、第1次充電（不確定 1）、第2次放電（短）",
          [st.dir for st in res.stages] == ["D", "C", "D"]
          and [len(st.pauses) for st in res.stages] == [1, 1, 0]
          and res.stages[0].pauses[0].certain and not res.stages[1].pauses[0].certain)
    check("頁籤順序：Summary／數據概覽／即時熱度／各階段（動態 3 個）／整體測試／控制紀錄／Raw Data",
          wb.sheetnames == ["Summary", "數據概覽", "即時熱度", "1.第1次放電", "2.第1次充電",
                            "3.第2次放電", "整體測試", "控制紀錄", "Raw Data"])
    wc = wb["控制紀錄"]
    check("來源沒有指令紀錄（離線資料）→ 控制紀錄頁顯示「無控制紀錄」，不推定歷史控制資訊",
          [wc.cell(1, c).value for c in range(1, 10)] == ["時間", "階段", "方向", "控制模式", "設定值", "單位",
                                                          "Action", "控制結果", "驗證結果"]
          and str(wc["A2"].value).startswith("無控制紀錄") and wc.max_row == 2)

    print("\n== 階段頁 ==")
    ws = wb["1.第1次放電"]
    COL = {h: i for i, h in enumerate(LOG_HEADERS, 1)}         # 欄名 → 欄號（不寫死欄位字母）
    NOTE = COL["備註"]
    check("欄位 A～K：時間、SOC、電流、Rack 最大／最小溫度、Cell 最大／最小電壓、總電壓、傳感器溫度、PCS 模組溫度、備註",
          [ws.cell(1, c).value for c in range(1, 12)] == ["時間", "SOC", "電流", "Rack 最大溫度", "Rack 最小溫度",
                                                         "Cell 最大電壓", "Cell 最小電壓", "總電壓", "傳感器溫度",
                                                         "PCS 模組溫度", "備註"] == LOG_HEADERS)
    times = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    check("每分鐘一列，只含階段範圍內的整分（10:00～10:20，21 列）",
          len(times) == 21 and times[0] == T0 and times[-1] == T0 + timedelta(minutes=20))
    check("時間欄格式 h:mm（不新增取樣時間欄）",
          ws.cell(2, 1).number_format == "h:mm" and ws.max_column <= 29)
    tbl = list(ws.tables.values())[0]
    check("Excel Table A1:K22（含備註欄）、樣式 TableStyleMedium7（同範本）",
          tbl.ref == "A1:K22" and tbl.tableStyleInfo.name == "TableStyleMedium7")
    cvx, cvn = ws.cell(2, COL["Cell 最大電壓"]), ws.cell(2, COL["Cell 最小電壓"])
    check(f"Cell 最大／最小電壓：原始精度保留（{cvx.value}／{cvn.value} V）、顯示格式 0.000、不是總電壓",
          cvx.value == 3.4015 and cvn.value == 3.3852 and cvx.number_format == "0.000"
          and cvn.number_format == "0.000" and ws.cell(2, COL["總電壓"]).value == 1200.0)
    check("Cell 電壓欄寬設定一致（F、G）",
          ws.column_dimensions["F"].width == ws.column_dimensions["G"].width >= 14)

    def row_of(ws, hhmm):
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, datetime) and v.strftime("%H:%M") == hhmm:
                return r
    r10 = row_of(ws, "10:10")
    check("停頓資料照實保留，階段頁備註保留分析標記並附設備狀態（無上下電資料 →「停頓（N/A）」）",
          str(ws.cell(r10, NOTE).value).startswith("停頓（N/A）")
          and ws.cell(r10, 3).value == 0)
    r05 = row_of(ws, "10:05")
    check("傳感器無效 → 顯示 N/A；圖表輔助欄為真正空白（不畫成 0）",
          ws.cell(r05, COL["傳感器溫度"]).value == "N/A" and ws.cell(r05, 29).value is None)
    r15 = row_of(ws, "10:15")
    check("有效 0 照實保留（SOC=0 為數值 0，不是 N/A）", ws.cell(r15, 2).value == 0)
    check("PCS 模組溫度無來源 → N/A（不虛構）", ws.cell(2, COL["PCS 模組溫度"]).value == "N/A")
    pair_ok = True
    for r in range(2, ws.max_row + 1):
        for vis, hid in ((COL["Rack 最大溫度"], 27), (COL["Rack 最小溫度"], 28), (COL["傳感器溫度"], 29)):
            v, h = ws.cell(r, vis).value, ws.cell(r, hid).value
            pair_ok &= (h is None) if v == "N/A" else (h == v)
    check("圖表輔助欄 AA／AB／AC 逐列等於 Rack 最大溫度（D）／Rack 最小溫度（E）／傳感器溫度（I）（N/A → 空白），"
          "不含電壓", pair_ok and [ws.cell(1, c).value for c in (27, 28, 29)]
          == ["（圖表用）Rack 最大溫度", "（圖表用）Rack 最小溫度", "（圖表用）傳感器溫度"])
    check("圖表輔助欄 AA～AC 隱藏",
          all(ws.column_dimensions[c].hidden for c in ("AA", "AB", "AC")))
    _anc = [c.anchor._from.col for c in wb["1.第1次放電"]._charts + wb["整體測試"]._charts]
    check(f"階段頁／整體頁圖表在備註欄（K）右側、不遮資料（起始欄 {[get_column_letter(c + 1) for c in _anc]}）",
          _anc and all(c >= NOTE for c in _anc))
    check("階段頁各 1 張圖表（短階段除外）",
          len(wb["1.第1次放電"]._charts) == 1 and len(wb["2.第1次充電"]._charts) == 1)
    ws3 = wb["3.第2次放電"]
    check("短階段：無整分鐘 Log、不畫圖，保留頭尾觀測值與說明",
          len(ws3._charts) == 0 and "無整分鐘 Log" in str(ws3["A1"].value)
          and "首筆觀測" in str(ws3["A4"].value))
    ws2 = wb["2.第1次充電"]
    r30 = row_of(ws2, "10:30")
    check("±20 秒內只有讀取失敗 → 整列 N/A 並寫明原因",
          ws2.cell(r30, 2).value == "N/A" and "讀取失敗" in str(ws2.cell(r30, NOTE).value))
    r35 = row_of(ws2, "10:35")
    check("舊資料沒有 Cell 電壓 → N/A（不以總電壓代替、不補值）",
          ws2.cell(r35, COL["Cell 最大電壓"]).value == "N/A" and ws2.cell(r35, COL["Cell 最小電壓"]).value == "N/A"
          and ws2.cell(r35, COL["總電壓"]).value == 1200.0)

    print("\n== 整體頁 ==")
    wa = wb["整體測試"]
    rr = row_of(wa, "10:22")
    check("整體頁休息期間：資料沒有保存電池上下電狀態 → 備註 N/A（不以電流 0／idle 推測為待機）",
          wa.cell(rr, NOTE).value == "N/A")
    check("整體頁同一欄位配置（含 Cell 電壓與 0.000 格式）",
          [wa.cell(1, c).value for c in range(1, 12)] == LOG_HEADERS
          and wa.cell(2, COL["Cell 最大電壓"]).number_format == "0.000")
    check("整體頁含全部時間（10:00～10:45）",
          wa.cell(2, 1).value == T0 and row_of(wa, "10:45") is not None)

    print("\n== Summary ==")
    ss = wb["Summary"]
    check("A1～A3 測試條件、A4「溫度曲線：」", ss["A4"].value == "溫度曲線：" and ss["A1"].value)
    anchors = [c.anchor._from.row for c in ss._charts]
    check("Summary 圖表 3 張（2 個階段 ＋ 整體；短階段以文字標示）",
          len(ss._charts) == 3 and any("不足一分鐘" in str(c.value) for row in ss.iter_rows()
                                       for c in row if c.value))
    below = [c for row in ss.iter_rows() for c in row
             if c.value is not None and c.row >= max(anchors) + 1]
    check("Summary 最後一張圖表下方沒有任何儲存格內容（無階段表、明細、參數、說明）", not below)
    texts = " ".join(str(c.value) for row in ss.iter_rows() for c in row if c.value)
    check("Summary 不含階段表／停頓明細／確認紀錄／名詞參數說明",
          not any(k in texts for k in ("階段表", "停頓與", "確認紀錄", "名詞與判定前提",
                                       "暫用（未定案）", "資料來源與限制")))
    item = res.stages[1].pauses[0].item_id
    check("箭頭不再用文字儲存格（→／←／↓ 間隔）",
          not any(k in texts for k in ("→", "←", "↓", "間隔")))

    with zipfile.ZipFile(path) as z:
        dr = z.read("xl/drawings/drawing1.xml").decode("utf-8")
    anchors_x = re.findall(r"<(?:xdr:)?twoCellAnchor[^>]*>(.*?)</(?:xdr:)?twoCellAnchor>", dr, re.S)

    def box(a):
        n = [int(v) for v in re.findall(r"<(?:xdr:)?(?:col|row)>(\d+)<", a)]
        return n[0], n[1], n[2], n[3]          # from col, from row, to col, to row
    shapes = [(re.search(r'prst="(\w+)"', a).group(1), box(a),
               re.findall(r"<a:t>([^<]*)</a:t>", a)) for a in anchors_x if "prstGeom" in a]
    frames = [box(a) for a in anchors_x if "graphicFrame" in a]
    check("Summary 有真正的箭頭圖形：第1→2 向右、第2→3 向下（蛇形）",
          [s[0] for s in shapes] == ["rightArrow", "downArrow"])
    check("箭頭內標示休息時間，由資料計算、只顯示整分鐘（R−B：300 秒 → 5 分、120 秒 → 2 分）",
          [s[2] for s in shapes] == [["休息", "5 分"], ["休息", "2 分"]])
    gap_obs = ([obs_at(s, "D") for s in range(0, 1201, 10)]
               + [obs_at(s, "C") for s in range(5249, 5550, 10)])      # A=10:20:00，R=11:27:29，中間無取樣
    rg = analyze(gap_obs, p)
    check("中間無資料但兩端已知 → 「間隔約 67 分／期間無資料」（R−A=4049 秒，捨去秒數，不稱休息）",
          arrow_text(rg, rg.stages[0], rg.stages[1]) == ["間隔約", "67 分", "期間無資料"])
    from cert_stage_engine import overall_minutes as _om
    gap_rows = [m for m in _om(rg) if T0 + timedelta(minutes=21) <= m.T <= T0 + timedelta(minutes=87)]
    check("同一缺口在全程資料仍為 N/A（箭頭有間隔不代表補值或連線）",
          gap_rows and all(m.obs is None for m in gap_rows))
    rf_obs = ([obs_at(s, "D") for s in range(0, 1201, 10)] + [obs_at(s, "I") for s in range(1210, 1501, 10)]
              + [obs_at(1510, "X")] + [obs_at(s, "I") for s in range(1520, 1801, 10)]
              + [obs_at(s, "C") for s in range(1810, 2101, 10)])
    rr_ = analyze(rf_obs, p)
    check("休息期間讀取失敗 → 「間隔約 10 分／期間讀取失敗」（R−A=610 秒）",
          arrow_text(rr_, rr_.stages[0], rr_.stages[1]) == ["間隔約", "10 分", "期間讀取失敗"])
    check("分鐘無條件捨去：800 秒（13 分 20 秒）→ 13 分、59 秒 → 0 分、3599 秒 → 59 分",
          [rest_minutes_text(x) for x in (800, 59, 3599)] == ["13 分", "0 分", "59 分"])

    def overlap(b1, b2):
        return not (b1[2] < b2[0] or b2[2] < b1[0] or b1[3] < b2[1] or b2[3] < b1[1])
    check("箭頭與任何圖表的儲存格範圍都不重疊",
          all(not overlap(s[1], f) for s in shapes for f in frames))
    def title_of(ch):
        return "".join(r.t for p_ in ch.title.tx.rich.p for r in (p_.r or []))
    check("Summary 圖表標題：第一次放電、第一次充電（未標記 → 用方向，不自動認定滿充／滿放）＋全程",
          [title_of(c) for c in ss._charts][:2] == ["第一次放電", "第一次充電"]
          and title_of(ss._charts[2]) == "【250KW】電櫃充放2個循環溫度曲線")
    check("整體溫度曲線標題：Summary 與「整體測試」頁一致＝「【250KW】電櫃充放2個循環溫度曲線」（規格取自測試資訊）",
          title_of(ss._charts[2]) == title_of(wb["整體測試"]._charts[0]) == "【250KW】電櫃充放2個循環溫度曲線")
    check("階段頁圖表標題共用同一電櫃規格（250KW 第1次放電）",
          title_of(wb["1.第1次放電"]._charts[0]) == "250KW 第1次放電")
    check("未輸入循環數 → 不寫循環數（不自動推算）；未輸入規格 → 【電櫃規格未設定】",
          overall_title({"cabinet_spec": "340KW"}) == "【340KW】電櫃充放溫度曲線"
          and overall_title({}) == "【電櫃規格未設定】電櫃充放溫度曲線"
          and overall_title({"cabinet_spec": "340KW", "cycles": 3}) == "【340KW】電櫃充放3個循環溫度曲線")
    from openpyxl.cell.rich_text import CellRichText
    ws_rt = load_workbook(path, rich_text=True)["Summary"]
    bold_ok = True
    for a, lab_ in (("A1", "冷氣溫度："), ("A2", "硬體概述："), ("A3", "測試條件：")):
        v = ws_rt[a].value
        bold_ok &= (isinstance(v, CellRichText) and str(v.__getitem__(0).text) == lab_
                    and v[0].font.b and not v[1].font.b)
    check("Summary 冷氣溫度：／硬體概述：／測試條件： 標籤粗體、內容一般字（同範本）", bold_ok)
    lab = [{"type": "stage_label", "stage_start": st.start.t.strftime("%Y-%m-%dT%H:%M:%S"),
            "label": w, "at": "2026-01-05 12:00:00", "by": "tester"}
           for st, w in zip(res.stages, ["滿放", "滿充", "滿放"])]
    check("有操作人員標記 → 標題「第一次滿放／第一次滿充／第二次滿放」（依標記各自計數）",
          list(summary_titles(analyze(scenario(), p, lab).stages).values())
          == ["第一次滿放", "第一次滿充", "第二次滿放"])
    check("中文序數：1→一、10→十、11→十一、23→二十三",
          [_cn_num(n) for n in (1, 10, 11, 23)] == ["一", "十", "十一", "二十三"])

    print("\n== 截圖頁 ==")
    for name in ("數據概覽", "即時熱度"):
        w = wb[name]
        ph = [c.value for row in w.iter_rows() for c in row if c.value and "待截圖" in str(c.value)]
        check(f"{name}：每個階段 1 個「待截圖」標示（共 3），沒有任何圖片",
              len(ph) == 3 and not w._images and all("不使用範本" in v for v in ph))
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        check("檔案內沒有任何圖片（xl/media）", not any(n.startswith("xl/media/") for n in names))
        charts = [n for n in names if re.match(r"xl/charts/chart\d+\.xml$", n)]
        xmls = [z.read(n).decode("utf-8") for n in charts]

    print("\n== 圖表 XML ==")
    check("圖表總數 2（階段頁，短階段無圖）＋1（整體頁）＋3（Summary）＝ 6", len(xmls) == 6)
    ok_all = True
    for x in xmls:
        x = re.sub(r"<(/?)c:", r"<\1", x)          # openpyxl 以預設命名空間輸出（無 c: 前綴）
        ok_all &= ('plotVisOnly val="0"' in x and 'dispBlanksAs val="gap"' in x
                   and x.count("<ser>") == 3 and "<lineChart>" in x
                   and re.search(r"\$D\$1.*\$E\$1.*\$I\$1", x, re.S) is not None
                   and "$F$" not in x and "$G$" not in x and "$H$" not in x
                   and re.search(r"\$AA\$2.*\$AB\$2.*\$AC\$2", x, re.S) is not None
                   and 'formatCode="h:mm"' in x and 'legendPos val="b"' in x
                   and '<majorUnit val="1"' in x and "<min " in x and "<max " in x
                   and re.search(r'<catAx>.*<axPos val="b"', x, re.S) is not None)
        gl = re.search(r"<majorGridlines>(.*?)</majorGridlines>", x, re.S)
        ok_all &= (gl is not None and 'schemeClr val="tx1"' in gl.group(1)
                   and 'lumMod val="15000"' in gl.group(1) and 'lumOff val="85000"' in gl.group(1))
        ymax = int(re.search(r'<max val="(\d+)"', x).group(1))
        ok_all &= ymax < 100                       # 溫度軸（防止誤接電壓等欄位）
    check("每張圖：折線、3 系列（名稱取 D1/E1/I1＝三條溫度，不引用 Cell 電壓 F/G、總電壓 H）、讀隱藏輔助欄、缺值斷線、plotVisOnly=0、"
          "h:mm、圖例在下、Y 軸手動範圍 majorUnit 1、水平格線淡灰（同範本 tx1 85%）", ok_all)

    print("\n== Raw Data 追溯 ==")
    rw = wb["Raw Data"]
    hdr = [rw.cell(1, c).value for c in range(1, 19)]
    rows = [[rw.cell(r, c).value for c in range(1, 19)] for r in range(2, rw.max_row + 1)]
    check("Raw Data：Rack 最小溫度後加入 Cell 最大／最小電壓，說明欄移到最後（R）",
          hdr[10:13] == ["Rack 最小溫度", "Cell 最大電壓", "Cell 最小電壓"] and hdr[17] == "說明"
          and rw.auto_filter.ref.startswith("A1:R"))
    _ok = [r for r in rows if r[3] == "成功" and r[11] not in (None, "N/A")]
    check(f"Raw Data Cell 電壓保留原始值（{_ok[0][11] if _ok else None}）",
          _ok and _ok[0][11] == 3.4015 and _ok[0][12] == 3.3852)
    obs_rows = [r for r in rows if r[3] in ("成功", "失敗")]
    na_rows = [r for r in rows if r[3] == "（無取樣）"]
    check("Raw Data 保留全部觀測（含讀取失敗）", len(obs_rows) == len(res.obs)
          and sum(1 for r in obs_rows if r[3] == "失敗") == 5)
    check("Raw Data 保留實際取樣時間（含秒）", hdr[1] == "PC 取樣時間"
          and obs_rows[0][1] == T0)
    m = [r for r in obs_rows if r[6] and "階段頁「1.第1次放電」 10:05（+0 秒）" in r[6]]
    check("階段頁每分鐘可對回 Raw Data 原始紀錄（10:05 → 10:05:00）",
          len(m) == 1 and m[0][1] == T0 + timedelta(minutes=5))
    check("N/A 分鐘在 Raw Data 有「（無取樣）」列與原因",
          any("10:30（N/A）" in str(r[6]) and "讀取失敗" in str(r[17]) for r in na_rows))

    print("\n== 操作人員確認與一致性 ==")
    dec = [{"item_id": item, "choice": "split", "at": "2026-01-05 12:10:00", "by": "tester"}]
    res2 = analyze(scenario(), p, dec)
    path2 = save_report(res2, META, os.path.join(tmp, "r_split.xlsx"))
    wb2 = load_workbook(path2)
    pz2 = [pz for pz in res2.pauses if pz.item_id == item]          # 拆分後不屬任何階段，但紀錄仍在
    check("選「拆分」→ 階段頁變為 4 個；底層紀錄仍保留系統原始判定與確認紀錄",
          len([n for n in wb2.sheetnames if re.match(r"\d+\.", n)]) == 4
          and len(pz2) == 1 and pz2[0].system_judgment.startswith("暫時保留同一階段")
          and pz2[0].decision and pz2[0].decision["choice"] == "split" and res2.decisions == dec)
    raw1 = [[c.value for c in row][:4] + [c.value for c in row][7:15]
            for row in wb["Raw Data"].iter_rows(min_row=2)]
    raw2 = [[c.value for c in row][:4] + [c.value for c in row][7:15]
            for row in wb2["Raw Data"].iter_rows(min_row=2)]
    check("確認後 Raw Data 的原始觀測欄位（時間、來源、讀取結果、數值）完全相同", raw1 == raw2)
    path3 = save_report(analyze(scenario(), p), META, os.path.join(tmp, "r_again.xlsx"))
    wb3 = load_workbook(path3)
    same = all([[c.value for c in row] for row in wb[n].iter_rows()] ==
               [[c.value for c in row] for row in wb3[n].iter_rows()] for n in wb.sheetnames)
    check("同樣輸入重複產生 → 所有頁籤內容相同", same and wb.sheetnames == wb3.sheetnames)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

ok = all(_results)
print(f"\n== 認證報告 P1 Excel 格式 {'PASS' if ok else 'FAIL'}（{sum(_results)}/{len(_results)} 檢查通過）==")
sys.exit(0 if ok else 1)
