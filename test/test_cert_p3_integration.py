# -*- coding: utf-8 -*-
"""
認證測試報告 P3 —— 服務／選單整合模擬測試（離線）
================================================
不連設備、不開網頁、不送任何控制、不碰正式服務與正式 Mutex：
  - 讀值：模擬設備（依時間表產生充電／放電／休息／通訊失敗／電流缺值）
  - 時間：模擬時鐘（每輪 +10 秒），可在數秒內跑完約 1.5 小時的認證
  - 截圖：假的截圖函式（複製 2026-10-27 確認成品，報告註明非本模擬畫面）
  - Ownership：output root 改到暫存資料夾 → Mutex 名稱與正式不同（同 test_auto_schedule 做法）
驗證：開始／結束／重複請求／逾時後查詢、單次↔認證雙向互斥、啟動順序與 Ownership、
      ≥5 分鐘休息持續取樣與全程曲線、通訊失敗、服務重啟續接、產報失敗重試、階段確認、
      截圖歸屬與不阻塞取樣、功能關閉時既有行為不變。
執行：python test_cert_p3_integration.py [輸出 Excel 副本資料夾]
"""

import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from datetime import datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import cert_report_config as CFG
import cert_requests as REQ
import report_monitor as RM
_DEPLOYED_CERT_MODE = CFG.CERT_MODE_ENABLED   # 現場設定值（部署後為 True）；測試不改檔案，只在本行程內切換
CFG.CERT_MODE_ENABLED = False                 # 以「功能關閉」匯入選單 → 檢查關閉時行為不變
import device_control_menu as DCM
import cert_service
import cert_menu
from cert_runtime import CertRuntime, _read_jsonl, record_to_obs
from cert_stage_engine import analyze
from cert_screenshot import ShotError, load_manifest, png_size, select_stage_shot, shot_params

_results = []


def check(label, ok):
    _results.append(bool(ok))
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}")


TMP = tempfile.mkdtemp(prefix="cert_p3_")
_BAK_ROOT = RM._report_output_root
RM._report_output_root = lambda: os.path.join(TMP, "cdr")      # 暫存 Mutex 名稱（不碰正式）
os.makedirs(RM._report_output_root(), exist_ok=True)
OUT_COPY = sys.argv[1] if len(sys.argv) > 1 else None
CROPS = os.path.join(HERE, "..", "output", "cert_reports_offline", "P2_layout_TW20251218_confirmed_crops", "screenshots")


# =============================================================================
# 模擬工具
# =============================================================================
class Clock:
    def __init__(self, t0):
        self.t = t0

    def __call__(self):
        return self.t

    def advance(self, sec):
        self.t += timedelta(seconds=sec)


class Plant:
    """依時間表（距 t0 秒數）產生讀值。mode：D 放電／C 充電／I 待機／FAIL 讀取失敗／NOCUR 電流功率缺值（溫度有值）。"""

    def __init__(self, clock, t0, plan):
        self.clock, self.t0, self.plan = clock, t0, plan
        self.temp = 25.0
        self.calls = 0

    def mode(self):
        s = (self.clock() - self.t0).total_seconds()
        for a, b, m in self.plan:
            if a <= s < b:
                return m
        return "I"

    def read(self):
        self.calls += 1
        m = self.mode()
        if m == "FAIL":
            return None
        self.temp += {"D": 0.012, "C": 0.010}.get(m, -0.004)
        cur = {"D": -140.0, "C": 135.0}.get(m, 0.0)
        r = {"soc_percent": 50, "battery_current_a": cur, "actual_active_power_kw": cur * 1.2,
             "rack_max_temperature_c": round(self.temp, 1), "rack_min_temperature_c": round(self.temp - 8, 1),
             "battery_voltage_v": 1250.0, "_fail": [], "communication_ok": True,
             "rack_cell_max_voltage_v": 3.4015, "rack_cell_min_voltage_v": 3.3852}
        r.update(battery_power_status="已上電", pcs_fault_flag=False, pcs_running_flag=m in ("C", "D"),
                 pcs_standby_flag=False, pcs_charging_flag=m == "C", pcs_discharging_flag=m == "D")
        if m == "OFF":                                  # 電池已下電（接觸器斷開）
            r.update(battery_power_status="已下電")
        if m == "FAULT":                                # PCS 故障、BCU 失聯 → 上下電未知
            r.update(battery_power_status="未知", pcs_fault_flag=True)
        if m == "NOCUR":
            r.update(battery_current_a=None, actual_active_power_kw=None, communication_ok=False,
                     _fail=["/hmiGuest/unauthorizedAccess/overview/mainControlCollectsInformation:Timeout"])
        return r

    def heat(self):
        if self.mode() == "FAIL":
            raise OSError("模擬：6161 無回應")
        return {"st": 22.5, "pt": {"moduleTemperature": 40}}


class FakeShot:
    """假截圖：複製確認成品；fail_calls＝第幾次 overview 呼叫失敗（驗證延遲補拍）；sleep 模擬耗時。"""

    def __init__(self, fail_calls=(), sleep=0.0):
        self.fail_calls, self.sleep, self.calls = set(fail_calls), sleep, {"overview": 0, "heat": 0}
        self.src = {k: sorted(glob.glob(os.path.join(CROPS, f"*_{k}.png")))[0] for k in ("overview", "heat")}

    def __call__(self, target, out_png, p, procs=None):
        self.calls[target.key] += 1
        time.sleep(self.sleep)
        if target.key == "overview" and self.calls["overview"] in self.fail_calls:
            raise ShotError("模擬：開頁失敗 net::ERR_CONNECTION_TIMED_OUT")
        shutil.copyfile(self.src[target.key], out_png)
        w, h = png_size(out_png)
        now = datetime.now()
        return {"captured_at": now, "rect": [0, 0, w, h], "size": [w, h],
                "data_time": now if target.key == "heat" else None, "page_clock": None, "notes": []}


def feed(answers):
    it = iter(answers)
    return lambda prompt="": next(it)


SP = shot_params(SHOT_ATTEMPTS=1, SHOT_RETRY_DELAY_SEC=0, SHOT_ATTEMPT_HARD_TIMEOUT_SEC=10)


def runtime(root, clock, plant, shots=None, single=None, cmd_file=None):
    rt = CertRuntime(root, reader=plant.read, heat_reader=plant.heat, single_state=single or (lambda: None),
                     cmd_file=cmd_file, shot_params=SP, capturer=shots or FakeShot(), enable_shots=True,
                     clock=clock, log=lambda *a: None)
    rt.extra_notes = ["本檔為 P3 模擬測試輸出（模擬設備讀值與模擬時鐘），不是實際認證測試報告。"]
    rt.shot_note = "模擬測試：截圖為 2026-10-27 範圍確認時的成品，非本模擬階段畫面。"
    return rt


def req(root, clock, rtype, payload=None, cert_id=None):
    return REQ.new_request(root, rtype, payload, cert_id=cert_id, now=clock())


def run_until(rt, clock, until_s, t0, step=10):
    while (clock() - t0).total_seconds() < until_s:
        clock.advance(step)
        rt.tick()


def idle_workers(*rts):
    for r in rts:
        if r.worker:
            r.worker.wait_idle(30)


try:
    # =========================================================================
    print("\n== 功能關閉（預設）：既有行為不變 ==")
    check(f"功能關閉（本行程強制 False；現場設定檔值 {_DEPLOYED_CERT_MODE}，未修改）", CFG.CERT_MODE_ENABLED is False)
    check("選單不顯示認證測試卡片、21～27 不處理", "認證測試" not in DCM.MENU_TEXT and not DCM._cert_enabled())
    check("服務端不啟動認證、不安裝互斥 hook",
          cert_service.start_if_enabled(None) is None and RM._cert_active_provider is None)
    check("RM 互斥 hook 未安裝時 _cert_block() 為 None（單次報告入口不受影響）", RM._cert_block() is None)
    import io
    import contextlib
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        DCM.handle_choice("21")
    check("功能關閉時輸入 21 → 仍是既有「無效選項」", "無效選項" in buf.getvalue())

    print("\n== read_all：Cell 最大／最小電壓取自同一份 rackExtremeValueInformation ==")
    import charge_discharge_report as CDR

    class _Cli:
        def __init__(self):
            self.paths = []

        def get(self, path, **kw):
            self.paths.append(path)
            if path == CDR.EP_RACKEXT:
                return {"rackCellMaxTemperature": "53", "rackCellMinTemperature": "42",
                        "rackCellMaxVoltage": "3.4015", "rackCellMinVoltage": 3.385,
                        "rackCellMaxVoltageIndex": "16", "rackCellMinVoltageIndex": "130"}
            if path == CDR.EP_MAINCTL:
                return {"rackSoc": 50, "rackTotalBatteryVoltage": 952.3, "rackElectricCurrent": -140.1}
            return None
    cli = _Cli()
    with contextlib.redirect_stdout(io.StringIO()):
        rd = CDR.read_all(cli)
    check(f"Cell 最大／最小電壓解析（{rd.get('rack_cell_max_voltage_v')}／{rd.get('rack_cell_min_voltage_v')} V），"
          "原始精度不降、不是總電壓",
          rd["rack_cell_max_voltage_v"] == 3.4015 and rd["rack_cell_min_voltage_v"] == 3.385
          and rd["battery_voltage_v"] == 952.3)
    check("不新增 API 請求：rackExtremeValueInformation 每次仍只請求 1 次",
          cli.paths.count(CDR.EP_RACKEXT) == 1)
    check("read_all 原有欄位與數值不變（Rack 溫度 53／42、SOC、電流）",
          rd["rack_max_temperature_c"] == 53.0 and rd["rack_min_temperature_c"] == 42.0
          and rd["soc_percent"] == 50.0 and rd["battery_current_a"] == -140.1)
    import inspect
    row_src = inspect.getsource(CDR.ReportSession._sample_row)
    check("單次報告 samples.csv 欄位不變（_sample_row 未引用 Cell 電壓欄位）",
          "rack_cell_max_voltage_v" not in row_src and "rack_cell_min_voltage_v" not in row_src)

    CFG.CERT_MODE_ENABLED = True
    t0 = datetime(2026, 11, 2, 9, 0, 0)

    # =========================================================================
    print("\n== 選單請求：等待逾時 → 結果待確認 → 查詢 ==")
    root = os.path.join(TMP, "cert_main")
    clock = Clock(t0)
    # 時間表（秒）：放電 0–1800（900–930 停頓 30 秒）→ 休息 6 分 → 充電 2160–3960（3000–3120 通訊失敗）
    #             → 休息 6 分（4140–4200 電流缺值、溫度有值）→ 放電 4320–5400（4800 服務重啟）→ 待機
    plan = [(0, 900, "D"), (900, 930, "I"), (930, 1800, "D"), (1800, 2160, "I"),
            (2160, 3000, "C"), (3000, 3120, "FAIL"), (3120, 3960, "C"),
            (3960, 4140, "I"), (4140, 4200, "NOCUR"), (4200, 4320, "I"),
            (4320, 5400, "D"), (5400, 99999, "I")]
    plant = Plant(clock, t0, plan)
    cmd_file = os.path.join(TMP, "device_control_action_result.json")
    REQ.atomic_write_json(cmd_file, {"timestamp": "2026-11-01 08:00:00", "action": "pcs_stop_power",
                                     "control_success": True, "dry_run": False})
    shots = FakeShot(fail_calls={2})      # 第 1 次＝停頓候選，第 2 次＝第 1 階段結束候選（失敗）
    rtA = runtime(root, clock, plant, shots, cmd_file=cmd_file)
    outs = []
    ui = cert_menu.UI(root, inp=feed(["340KW", "2", "製冷21℃", "潤泰退回pack", "定功率170kw", "YES"]),
                      out=outs.append, now=clock, wait_sec=0.3)
    REQ.atomic_write_json(os.path.join(root, "status.json"),
                          {"updated_at": REQ.now_str(clock()), "pid": 1, "owner_role": "service"})
    t_ui = time.monotonic()
    r_start = ui.start()
    check(f"服務尚未處理 → 選單等待 {time.monotonic() - t_ui:.1f} 秒後顯示「結果待確認」，不自動重送",
          r_start is None and any("結果待確認" in o for o in outs)
          and len(REQ.pending_requests(root)) == 1)
    rid_start = REQ.pending_requests(root)[0]["request_id"]
    check("逾時後查詢：狀態為「待處理」", REQ.request_state(root, rid_start)[0] == "pending")
    rtA.tick()
    st, res = REQ.request_state(root, rid_start)
    cid = res["cert_id"] if res else None
    check(f"服務處理後查詢：完成，所屬認證 {cid}", st == "done" and cid and cid.startswith("CERT_"))
    outs.clear()
    ui2 = cert_menu.UI(root, inp=feed([rid_start]), out=outs.append, now=clock)
    ui2.query()
    check("選單「27 查詢請求結果」顯示完成與所屬認證", any("完成" in o and cid in o for o in outs))
    state = REQ.read_json(os.path.join(root, cid, "cert_state.json"))
    check("測試資訊照輸入記錄（電櫃規格 340KW 由輸入而來、循環數 2）",
          state["test_info"]["cabinet_spec"] == "340KW" and state["test_info"]["cycles"] == 2)

    print("\n== 重複請求保護 ==")
    shutil.copyfile(os.path.join(root, "requests", "done", f"{rid_start}.json"),
                    os.path.join(root, "requests", "pending", f"{rid_start}.json"))
    before = REQ.read_result(root, rid_start)
    rtA.tick()
    check("同一請求 ID 再次出現 → 不再執行（結果不變、沒有第二份認證）",
          REQ.read_result(root, rid_start) == before and len(glob.glob(os.path.join(root, "CERT_*"))) == 1)
    rid2 = req(root, clock, "start_cert", {"cabinet_spec": "250KW"})
    rtA.tick()
    r2 = REQ.read_result(root, rid2)
    check(f"認證進行中再送「開始」→ 拒絕（{r2['message'][:30]}…）",
          r2["status"] == "rejected" and cid in r2["message"] and len(glob.glob(os.path.join(root, "CERT_*"))) == 1)
    old = REQ.new_request(root, "end_cert", cert_id=cid, now=clock() - timedelta(seconds=CFG.CERT_REQUEST_EXPIRE_SEC + 5))
    rtA.tick()
    check("逾期請求（服務未執行時建立、太晚才處理）→ expired，不執行（認證仍在記錄）",
          REQ.read_result(root, old)["status"] == "expired" and rtA.is_active())
    bad = req(root, clock, "manual_shot", cert_id="CERT_19990101_000000")
    rtA.tick()
    check("所屬認證 ID 不符 → 拒絕", REQ.read_result(root, bad)["status"] == "rejected")

    print("\n== 單次報告 ↔ 認證 雙向互斥 ==")
    owned, why = RM.acquire_ownership(role="service")
    check(f"（測試）以暫存 Mutex 取得 Ownership：{why}", owned)
    RM.set_cert_active_provider(rtA.active_id)
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        s1 = RM._report_start_session("charge", "交流有功", origin="manual")
        RM.report_start()
        RM.report_resume_on_launch()
        act = RM.auto_schedule_check({"pcs_control_mode_code": "smart", "pcs_schedule_enabled": True,
                                      "pcs_charging_flag": True, "pcs_discharging_flag": False})
    check("認證進行中：_report_start_session（手動跟隨／排程）拒絕建立單次報告", s1 == (None, False))
    check(f"認證進行中：auto_schedule_check → {act}", act == ("skipped", "cert_active"))
    check("認證進行中：report_start／resume 皆拒絕並說明",
          buf.getvalue().count("認證測試") >= 3 and RM._report["session"] is None
          and not glob.glob(os.path.join(RM._report_output_root(), "*_auto")))
    RM.set_cert_active_provider(lambda: (_ for _ in ()).throw(RuntimeError("x")))
    with contextlib.redirect_stdout(io.StringIO()):
        check("互斥查詢例外 → fail closed（視為認證中，不建單次報告）", RM._cert_block() == "unknown")
    RM.set_cert_active_provider(None)
    root_b = os.path.join(TMP, "cert_single")
    clk_b = Clock(t0)
    RM._report["session"] = type("S", (), {"session_id": "20261102_090000_auto"})()
    rtS = CertRuntime(root_b, reader=Plant(clk_b, t0, []).read, single_state=cert_service.single_report_state,
                      enable_shots=False, clock=clk_b, log=lambda *a: None)
    rq = req(root_b, clk_b, "start_cert", {"cabinet_spec": "250KW"})
    rtS.tick()
    rres = REQ.read_result(root_b, rq)
    check(f"單次報告記錄中 → 拒絕開始認證並顯示狀態（{rres['message'][:34]}…）",
          rres["status"] == "rejected" and "20261102_090000_auto" in rres["message"]
          and not glob.glob(os.path.join(root_b, "CERT_*")))
    RM._report["session"] = None
    RM.release_ownership()

    print("\n== 持續取樣：放電 → ≥5 分鐘休息 → 充電（含通訊失敗）→ 休息（電流缺值）→ 放電 ==")
    t_s = time.monotonic()
    REQ.atomic_write_json(cmd_file, {"timestamp": "2026-11-02 09:30:00", "action": "pcs_stop_power",
                                     "control_success": True, "dry_run": False})
    run_until(rtA, clock, 4800, t0)
    n_before = len(rtA.records)
    check(f"開始到 4800 秒每 10 秒一筆，不因階段結束或休息停止（{n_before} 筆，{time.monotonic() - t_s:.1f} 秒）",
          n_before == 481)
    recs = rtA.records
    rest = [r for r in recs if "09:30:30" <= r["t"][11:] <= "09:35:50"]
    check(f"休息 6 分鐘期間持續取樣且溫度有值（{len(rest)} 筆）",
          len(rest) == 33 and all(r["ok"] and isinstance(r["values"]["rack_max_c"], float) for r in rest))
    fail = [r for r in recs if "09:50:00" <= r["t"][11:] < "09:52:00"]
    check("通訊失敗 2 分鐘：照寫一筆 ok=False＋原因（不沿用上一筆、不補值）",
          len(fail) == 12 and all(not r["ok"] and r["values"]["rack_max_c"] is None and "HMI" in r["fail_reason"]
                                  for r in fail))
    nocur = [r for r in recs if "10:09:00" <= r["t"][11:] < "10:10:00"]
    check("電流缺值但溫度有值：整筆保留（ok=True），溫度照存、電流記缺值",
          nocur and all(r["ok"] and r["values"]["current_a"] is None and r["values"]["rack_max_c"] is not None
                        and "current_a" in r["missing"] for r in nocur))
    cmds = _read_jsonl(os.path.join(root, cid, "commands.jsonl"))
    check("認證取樣保存 Cell 最大／最小電壓（原始值）",
          recs[0]["values"]["cell_vmax_v"] == 3.4015 and recs[0]["values"]["cell_vmin_v"] == 3.3852)
    check("控制結果檔：開始前的舊指令不記；之後的新指令以快照記入（註明非完整歷史）",
          len(cmds) == 1 and cmds[0]["t"] == "2026-11-02T09:30:00" and "只保留最近一次" in cmds[0]["note"])

    print("\n== 服務重啟續接（停機 3 分鐘）==")
    idle_workers(rtA)
    jobs_before = [r for r in load_manifest(os.path.join(root, cid)) if r["type"] == "job"]
    rtA.worker._append({"type": "job", "job_id": "20261102_100000_manual_dead00", "kind": "manual",
                        "stage_no": 3, "stage_start": None, "stage_name": "x", "reason": "模擬：重啟前未完成",
                        "trigger_time": None, "requested_at": "2026-11-02T10:20:00",
                        "targets": ["overview", "heat"], "ref_job": None, "extra": {}})
    del rtA                                             # 模擬服務停止（不經結束認證）
    clock.advance(180)
    rtB = runtime(root, clock, plant, FakeShot(), cmd_file=cmd_file)
    resumed = rtB.startup()
    st = REQ.read_json(os.path.join(root, cid, "cert_state.json"))
    ev = _read_jsonl(os.path.join(root, cid, "events.jsonl"))
    check(f"重啟後接續同一份認證 {resumed}（不建立新認證），記錄中斷期間與原因",
          resumed == cid and len(glob.glob(os.path.join(root, "CERT_*"))) == 1
          and st["interruptions"] and 180 <= st["interruptions"][-1]["seconds"] <= 200
          and any(e["type"] == "resume" for e in ev))
    m = load_manifest(os.path.join(root, cid))
    check("重啟前未完成的截圖工作 → 補記失敗（服務中斷），不會永遠顯示「尚未完成」",
          sum(1 for r in m if r["type"] == "result" and r["job_id"] == "20261102_100000_manual_dead00"
              and "服務中斷" in r["reason"]) == 2)
    run_until(rtB, clock, 5520, t0)
    check("續接後在同一資料夾持續取樣", len(rtB.records) > n_before + 50
          and len(_read_jsonl(os.path.join(root, cid, "observations.jsonl"))) == len(rtB.records))

    print("\n== 截圖：不阻塞取樣、歸屬與類型 ==")
    slow = FakeShot(sleep=1.5)
    rtB.worker.capturer = slow
    rtB.worker.submit("manual", 9, None, "測試", "不阻塞測試")
    time.sleep(0.2)
    t1 = time.monotonic()
    clock.advance(10)
    rtB.sample_once()
    dt = time.monotonic() - t1
    check(f"截圖執行中（每張 1.5 秒）取樣仍立即完成（{dt * 1000:.0f} ms）", dt < 0.5)
    rtB.worker.wait_idle(30)
    rtB.worker.capturer = FakeShot()
    idle_workers(rtB)
    res = rtB.analysis()
    keys = [st_.start.t.strftime("%Y-%m-%dT%H:%M:%S") for st_ in res.stages]
    m = load_manifest(os.path.join(root, cid))
    jobs = {r["job_id"]: r for r in m if r["type"] == "job"}
    resolves = [r for r in m if r["type"] == "resolve"]
    check(f"階段：{len(res.stages)} 個（放電／充電／放電），休息不切斷整份認證",
          [s.dir for s in res.stages] == ["D", "C", "D"])
    pz = [r for r in resolves if r["as"] == "pause"]
    check("放電中 30 秒停頓 → 候選截圖判為停頓點（只留資料夾）",
          len(pz) >= 1 and all(jobs[r["job_id"]]["stage_start"] == keys[0] for r in pz[:1]))
    se = [r for r in resolves if r["as"] == "stage_end"]
    check("休息達 60 秒 → 候選截圖判為階段結束，歸屬到對應階段開始時間",
          {jobs[r["job_id"]]["stage_start"] for r in se} >= {keys[0], keys[1]})
    dl = [j for j in jobs.values() if j["kind"] == "delayed" and j["targets"] == ["overview"]]
    check("第 1 階段數據概覽截圖失敗 → 對該頁延遲補拍（報告標示非停止當下畫面）",
          dl and dl[0]["stage_start"] == keys[0])
    j1, r1 = select_stage_shot(m, keys[0], "overview")
    j1h, r1h = select_stage_shot(m, keys[0], "heat")
    check("報告選圖：第 1 階段 數據概覽＝延遲補拍、即時熱度＝階段結束截圖（兩頁分別記錄）",
          j1["effective_kind"] == "delayed" and r1["status"] == "ok" and j1h["effective_kind"] == "stage_end")

    print("\n== 階段確認、結束認證（設備待機）、報告 ==")
    unc = [p for p in res.pauses if not p.certain]
    check(f"通訊失敗與服務重啟造成「階段是否連續不確定」項目：{[p.item_id for p in unc]}", len(unc) >= 2)
    rq_l = req(root, clock, "stage_label", {"stage_start": keys[0], "label": "滿放"}, cert_id=cid)
    rq_x = req(root, clock, "stage_label", {"stage_start": keys[0], "label": "很滿"}, cert_id=cid)
    rq_c = req(root, clock, "confirm_item", {"item_id": unc[0].item_id, "choice": "merge", "note": "現場確認"}, cert_id=cid)
    rtB.tick()
    check("階段標記（滿放）與不確定項目確認（同一階段）完成；不合法標記拒絕",
          REQ.read_result(root, rq_l)["status"] == "done" and REQ.read_result(root, rq_x)["status"] == "rejected"
          and REQ.read_result(root, rq_c)["status"] == "done")
    dec = _read_jsonl(os.path.join(root, cid, "decisions.jsonl"))
    check("操作人員確認紀錄 append-only 保留（含請求 ID、時間）；原始資料不變",
          len(dec) == 2 and all(d.get("request_id") and d.get("at") for d in dec))
    rq_m = req(root, clock, "manual_shot", {"reason": "現場補拍"}, cert_id=cid)
    rtB.tick()
    check("立即截圖／手動補拍：排入背景，歸屬目前階段",
          REQ.read_result(root, rq_m)["status"] == "done"
          and REQ.read_result(root, rq_m)["data"]["stage_start"] == keys[-1])
    idle_workers(rtB)
    outs = []
    rtB.write_status()
    cert_menu.UI(root, out=outs.append, now=clock).status()
    check("選單「26 查看目前認證狀態」：認證 ID、取樣數、階段、不確定項目（唯讀狀態檔）",
          any(cid in o and "取樣" in o for o in outs) and any(o.strip().startswith("階段 1.") for o in outs)
          and any("不確定" in o for o in outs))
    outs = []
    r_ = cert_menu.UI(root, inp=feed(["no"]), out=outs.append, now=clock, wait_sec=0.2).end()
    check("選單「22 結束認證」：提示不會停止設備；未輸入 YES → 取消、不送請求",
          r_ is None and any("不會" in o and "停止設備" in o for o in outs) and not REQ.pending_requests(root))
    rq_e = req(root, clock, "end_cert", cert_id=cid)
    rtB.tick()
    re_ = REQ.read_result(root, rq_e)
    st = REQ.read_json(os.path.join(root, cid, "cert_state.json"))
    check(f"結束認證（設備已待機）：只結束記錄、產生報告；未標記未完成（{re_['message'][:40]}…）",
          re_["status"] == "done" and st["status"] == "ended" and not st["incomplete"]
          and re_["data"]["report_ok"] and not REQ.active_on_disk(root))
    n_end = len(_read_jsonl(os.path.join(root, cid, "observations.jsonl")))
    clock.advance(60)
    rtB.tick()
    check("結束後不再取樣", len(_read_jsonl(os.path.join(root, cid, "observations.jsonl"))) == n_end)
    report1 = os.path.join(root, cid, re_["data"]["report"]) if not os.path.isabs(re_["data"]["report"]) \
        else re_["data"]["report"]

    from openpyxl import load_workbook
    wb = load_workbook(report1)
    ws = wb["整體測試"]
    rows = {c.value.strftime("%H:%M"): r for r, c in enumerate(ws["A"][1:], 2) if hasattr(c.value, "strftime")}

    def val(hhmm, col=4):
        r = rows.get(hhmm)
        return ws.cell(r, col).value if r else "（無此列）"
    rest_vals = [val(f"09:{m:02d}") for m in range(31, 36)]
    check(f"「整體測試」休息期間每分鐘都有溫度（09:31～09:35：{rest_vals}）",
          all(isinstance(v, (int, float)) for v in rest_vals))
    # 失敗 09:50:00～09:52:00：09:50 取 ±20 秒內的 09:49:50（有效讀值，依規則）；09:51 前後 20 秒全是失敗 → N/A
    fail_vals = [val("09:51")]
    check(f"通訊失敗分鐘（09:51，±20 秒內無有效讀值）為 N/A、曲線斷線（{fail_vals}），不補值",
          all(v in (None, "N/A") for v in fail_vals) and isinstance(val("09:50"), (int, float)))
    gap_vals = [val("10:21"), val("10:22")]
    check(f"服務停機期間（10:20～10:23）為 N/A（{gap_vals}）", all(v in (None, "N/A") for v in gap_vals))
    check(f"報告「整體測試」Cell 最大／最小電壓（F／G）：{val('09:10', 6)}／{val('09:10', 7)}，格式 0.000",
          val("09:10", 6) == 3.4015 and val("09:10", 7) == 3.3852
          and ws.cell(rows["09:10"], 6).number_format == "0.000")
    check(f"電流缺值的分鐘溫度仍保留（10:09 Rack 最大 {val('10:09')}，電流 {val('10:09', 3)}）",
          isinstance(val("10:09"), (int, float)) and val("10:09", 3) in (None, "N/A"))
    z = zipfile.ZipFile(report1)
    charts = [z.read(n).decode() for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n)]
    overall = [c for c in charts if "'整體測試'!" in c]
    refs = {tuple(re.findall(r"<(?:c:)?f>([^<]+)</(?:c:)?f>", c)) for c in overall}
    check("Summary 全程曲線與「整體測試」圖表引用同一資料範圍（同一份全程資料）", len(overall) == 2 and len(refs) == 1)
    titles_ = " ".join(re.findall(r"<a:t>([^<]*)</a:t>", " ".join(charts)))
    desc = wb.properties.description or ""
    check("報告標題用輸入的電櫃規格與循環數（【340KW】電櫃充放2個循環溫度曲線）",
          "【340KW】電櫃充放2個循環溫度曲線" in titles_)
    check("來源說明（暫用參數、指令快照限制、服務中斷期間）寫入活頁簿屬性，不改動已確認版面",
          "暫用測試值" in desc and "只保留最近一次" in desc and "服務中斷" in desc and "Summary" in wb.sheetnames[0])
    check("階段頁籤依操作人員標記命名（1.第1次放電(滿放)）", any(n.startswith("1.第1次放電(滿放)") for n in wb.sheetnames))
    z.close()

    print("\n== 產報失敗 → 資料保留 → 重產（不覆蓋）==")
    import cert_report_xlsx
    real = cert_report_xlsx.save_report

    def boom(*a, **k):
        open(a[2], "wb").write(b"partial")
        raise PermissionError("模擬：檔案被 Excel 開啟鎖住")
    cert_report_xlsx.save_report = boom
    rq_g = req(root, clock, "regenerate", cert_id=cid)
    rtB.tick()
    rg = REQ.read_result(root, rq_g)
    st = REQ.read_json(os.path.join(root, cid, "cert_state.json"))
    check("產報失敗 → 結果 failed、記錄錯誤、資料與舊報告保留、不留下半份檔案",
          rg["status"] == "failed" and st["last_report_error"] and os.path.exists(report1)
          and not glob.glob(os.path.join(root, cid, "reports", "*.partial.xlsx"))
          and len(_read_jsonl(os.path.join(root, cid, "observations.jsonl"))) == n_end)
    cert_report_xlsx.save_report = real
    rq_g2 = req(root, clock, "regenerate", cert_id=cid)
    rtB.tick()
    rg2 = REQ.read_result(root, rq_g2)
    reps = sorted(glob.glob(os.path.join(root, cid, "reports", "*.xlsx")))
    check(f"重產成功 → 新檔（{[os.path.basename(x) for x in reps]}），不覆蓋既有報告",
          rg2["status"] == "done" and len(reps) == 2 and reps[0] == report1)

    print("\n== 結束認證時設備仍運行 → 提示並標記未完成（不送停止）==")
    root_r = os.path.join(TMP, "cert_running")
    clk_r = Clock(t0)
    pl_r = Plant(clk_r, t0, [(0, 99999, "C")])
    rtR = runtime(root_r, clk_r, pl_r)
    rqs = req(root_r, clk_r, "start_cert", {"cabinet_spec": "250KW"})
    rtR.tick()
    cid_r = REQ.read_result(root_r, rqs)["cert_id"]
    run_until(rtR, clk_r, 300, t0)
    rqe = req(root_r, clk_r, "end_cert", cert_id=cid_r)
    rtR.tick()
    rr = REQ.read_result(root_r, rqe)
    st_r = REQ.read_json(os.path.join(root_r, cid_r, "cert_state.json"))
    check(f"仍在充電 → 結果提示、狀態標記未完成（{st_r['incomplete']}）",
          "仍在充電" in rr["message"] and st_r["incomplete"] and "未送出停止指令" in st_r["incomplete"])
    import cert_runtime as _CR
    src = open(_CR.__file__, encoding="utf-8").read() + open(cert_service.__file__, encoding="utf-8").read()
    check("認證程式不含任何控制出口（operator／execute／POST）",
          not any(k in src for k in ("device_control_operator", "--execute", "method=\"POST\"", ".post(")))
    wb_r = load_workbook(rr["data"]["report"])
    check(f"未完成報告：檔名標示（{os.path.basename(rr['data']['report'])}）、屬性描述註明未完成",
          rr["data"]["report"].endswith("_未完成.xlsx") and "未完成" in (wb_r.properties.description or ""))

    print("\n== 啟動順序與 Ownership（暫存 Mutex，與正式服務無關）==")
    CHILD = ("import sys,os;sys.path.insert(0,{here!r});import report_monitor as RM;"
             "RM._report_output_root=lambda:{root!r};ok,why=RM.acquire_ownership(role={role!r});"
             "print('OWNED' if ok else 'NOT:'+why,flush=True);sys.stdin.readline();RM.release_ownership()")

    def child(role):
        p = subprocess.Popen([sys.executable, "-c", CHILD.format(here=HERE, root=RM._report_output_root(), role=role)],
                             stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
        line = ""
        while not (line.startswith("OWNED") or line.startswith("NOT:")):
            line = p.stdout.readline()
            if not line:
                break
            line = line.strip()
        return p, line

    def child_end(p):
        try:
            p.stdin.write("\n")
            p.stdin.flush()
            p.wait(10)
        except Exception:
            p.kill()
    root_o = os.path.join(TMP, "cert_owner")
    os.makedirs(root_o, exist_ok=True)
    # ① 服務先啟動 → 選單為 Observer，可送請求
    p, line = child("service")
    owned, why, ca = DCM._dashboard_acquire(cert_root=root_o)
    check(f"① 服務先啟動（{line}）→ 選單 Observer（{why}），可送認證請求",
          line == "OWNED" and not owned and why == "held_by_other"
          and cert_menu.UI(root_o, is_owner=RM.is_owner).is_owner() is False)
    child_end(p)
    # ② 選單先啟動（無認證）→ 選單 Owner；服務起不來；選單拒絕開始認證（必須由服務記錄）
    owned, why, ca = DCM._dashboard_acquire(cert_root=root_o)
    p, line = child("service")
    outs = []
    r_ = cert_menu.UI(root_o, inp=lambda *_: "YES", out=outs.append, is_owner=RM.is_owner).start()
    check(f"② 選單先啟動 → 選單 Owner（{why}）、服務 {line}；選單拒絕開始認證並說明原因",
          owned and line == "NOT:held_by_other" and r_ is None and any("背景服務無法接手" in o for o in outs))
    check("② 選單（dashboard Owner）不會啟動認證記錄", cert_service.start_if_enabled(None, log=lambda *a: None) is None)
    child_end(p)
    RM.release_ownership()
    # ③ 認證進行中、服務暫停 → 開啟選單**不取得** Ownership → 服務可接手並續接
    REQ.atomic_write_json(os.path.join(root_o, "active.json"), {"cert_id": "CERT_X"})
    os.makedirs(os.path.join(root_o, "CERT_X"), exist_ok=True)
    REQ.atomic_write_json(os.path.join(root_o, "CERT_X", "cert_state.json"), {"cert_id": "CERT_X", "status": "recording"})
    owned, why, ca = DCM._dashboard_acquire(cert_root=root_o)
    p, line = child("service")
    check(f"③ 認證進行中開啟選單 → 不取得 Ownership（{why}），服務可接手（{line}）",
          not owned and ca == "CERT_X" and line == "OWNED" and not RM.is_owner())
    child_end(p)
    owned, why = RM.acquire_ownership(role="service")
    import threading
    ev = threading.Event()
    _bak_root_fn = REQ.cert_root
    REQ.cert_root = lambda output_dir=None: root_o
    _bak_rd, _bak_hr = cert_service._reader, cert_service._heat_reader
    cert_service._reader = lambda: None
    cert_service._heat_reader = lambda: (_ for _ in ()).throw(OSError("測試：不連 6161"))   # 測試不碰網路
    rt_svc = cert_service.start_if_enabled(ev, log=lambda *a: None)
    check("③ 服務（service Owner）啟動 → 接續認證並安裝互斥",
          rt_svc is not None and rt_svc.active_id() == "CERT_X" and RM._cert_block() == "CERT_X")
    time.sleep(2.5)
    sts = REQ.read_status(root_o) or {}
    check(f"③ 正式模式（兩條執行緒）：無 client 時照寫讀取失敗紀錄（{len(rt_svc.records)} 筆）、狀態檔每秒更新",
          rt_svc.records and not rt_svc.records[0]["ok"] and sts.get("cert_id") == "CERT_X"
          and sts.get("owner_role") == "service")
    ev.set()
    time.sleep(1.2)
    cert_service._reader, cert_service._heat_reader = _bak_rd, _bak_hr
    REQ.cert_root = _bak_root_fn
    RM.set_cert_active_provider(None)
    RM.release_ownership()

    print("\n== 控制紀錄可追溯性：第1次放電＝直流恆流 100 A、第2次放電＝直流恆功率 100 kW ==")
    import cert_controls
    import device_control_operator as OPR                 # 只讀 ACTIONS 對照（不呼叫 run、不送指令）
    pw = {k: v for k, v in OPR.ACTIONS.items() if v.get("kind") == "pcs_power"}
    check(f"集中對照表涵蓋既有全部 {len(pw)} 個 PCS 充放電／停止 action，方向與模式和 operator 定義一致（不改 action 名稱）",
          set(pw) == set(cert_controls.ACTION_INFO)
          and all(cert_controls.ACTION_INFO[a][0] == (v.get("direction") or "stop")
                  and cert_controls.ACTION_INFO[a][1] == (v.get("pcs_mode") or "stop") for a, v in pw.items()))
    root_c = os.path.join(TMP, "cert_ctrl")
    clk_c = Clock(t0)
    # 時間表：待機 60s → 放電（直流恆流）→ 停 → 待機 3 分 → 放電（直流恆功率）→ 停 → 待機
    pl_c = Plant(clk_c, t0, [(0, 60, "I"), (60, 600, "D"), (600, 780, "I"), (780, 1380, "D"), (1380, 99999, "I")])
    cf = os.path.join(TMP, "ctrl_result.json")
    rtC = runtime(root_c, clk_c, pl_c, cmd_file=cf)
    rqs = req(root_c, clk_c, "start_cert", {"cabinet_spec": "250KW"})
    rtC.tick()
    cid_c = REQ.read_result(root_c, rqs)["cert_id"]
    script = {55: ("pcs_dc_current_discharge", 100.0), 605: ("pcs_stop_power", None),
              775: ("pcs_dc_power_discharge", 100.0), 1385: ("pcs_stop_power", None)}
    while (clk_c() - t0).total_seconds() < 1560:
        clk_c.advance(5)
        s_ = int((clk_c() - t0).total_seconds())
        if s_ in script:
            act, pv = script[s_]
            REQ.atomic_write_json(cf, {"timestamp": clk_c().strftime("%Y-%m-%d %H:%M:%S"), "action": act,
                                       "power": pv, "dry_run": False, "control_success": True,
                                       "verify_success": True, "success": True, "blocked": False,
                                       "control_request": {"payload": {"x": 1}}})
        rtC.tick()
    cmds = _read_jsonl(os.path.join(root_c, cid_c, "commands.jsonl"))
    check(f"commands.jsonl 四筆依序新增、不互相覆蓋（{[c['action'] for c in cmds]}）",
          [c["action"] for c in cmds] == ["pcs_dc_current_discharge", "pcs_stop_power",
                                          "pcs_dc_power_discharge", "pcs_stop_power"])
    check("100 A／100 kW 設定值、單位、控制模式、控制／驗證／整體結果都保存",
          (cmds[0]["setpoint"], cmds[0]["unit"], cmds[0]["control_mode"]) == (100.0, "A", "dc_current")
          and (cmds[2]["setpoint"], cmds[2]["unit"], cmds[2]["control_mode"]) == (100.0, "kW", "dc_power")
          and cmds[1]["control_mode"] == "stop" and cmds[1]["setpoint"] is None and cmds[1]["unit"] is None
          and all(c["control_success"] is True and c["verify_success"] is True and c["success"] is True for c in cmds))
    check("原始指令不寫入階段（階段於產報時對應）、不含 token／payload",
          all("stage" not in c and "stage_no" not in c and "control_request" not in c for c in cmds))
    rqe = req(root_c, clk_c, "end_cert", cert_id=cid_c)
    rtC.tick()
    rep_c = REQ.read_result(root_c, rqe)["data"]["report"]
    wbc = load_workbook(rep_c)
    wsc = wbc["控制紀錄"]
    crow = [[wsc.cell(r, c).value for c in range(1, 12)] for r in range(2, wsc.max_row + 1)]
    for row_ in crow:
        print("     ", row_[0], "|", row_[1], "|", row_[2], "|", row_[3], "|", row_[4], row_[5], "|", row_[6], "|", row_[7], row_[8], row_[9])
    check("Excel「控制紀錄」：第1次放電＝直流恆流 100 A、第2次放電＝直流恆功率 100 kW，停止指令各自歸屬",
          [(r_[1], r_[2], r_[3], r_[4], r_[5]) for r_ in crow]
          == [("第1次放電", "放電", "直流恆流", 100, "A"), ("第1次放電", "停止", "停止", "N/A", "N/A"),
              ("第2次放電", "放電", "直流恆功率", 100, "kW"), ("第2次放電", "停止", "停止", "N/A", "N/A")]
          and all(r_[7] == "成功" and r_[8] == "成功" for r_ in crow))
    check("頁籤：Summary／數據概覽／即時熱度 不變，控制紀錄在整體測試之後、Raw Data 之前",
          wbc.sheetnames[:3] == ["Summary", "數據概覽", "即時熱度"]
          and wbc.sheetnames[-3:] == ["整體測試", "控制紀錄", "Raw Data"])

    print("\n== 控制紀錄：舊資料相容、無法對應階段 → N/A ==")
    cd_ = os.path.join(root_c, cid_c)
    old = [{"t": c["t"], "cmd": c["cmd"], "action": c["action"], "control_success": c["control_success"],
            "dry_run": False, "seen_at": c["seen_at"], "note": c["note"]} for c in cmds]   # P3 初版欄位
    old.append({"t": (t0 + timedelta(seconds=700)).strftime("%Y-%m-%dT%H:%M:%S"), "cmd": "charge",
                "action": "pcs_dc_current_charge", "control_success": True, "dry_run": False})  # 休息中、無充電階段
    old.append({"t": (t0 + timedelta(seconds=700)).strftime("%Y-%m-%dT%H:%M:%S"), "cmd": None,
                "action": "ac_on", "control_success": True, "dry_run": False})                  # 非充放電控制
    with open(os.path.join(cd_, "commands.jsonl"), "w", encoding="utf-8") as f:
        for c in old:
            f.write(json.dumps(c, ensure_ascii=False) + "\n")
    rqg = req(root_c, clk_c, "regenerate", cert_id=cid_c)
    rtC.tick()
    rg_ = REQ.read_result(root_c, rqg)
    check(f"舊格式 commands.jsonl（缺新欄位）→ 報告照常產生、無 Exception（{rg_['status']}）", rg_["status"] == "done")
    wso = load_workbook(rg_["data"]["report"])["控制紀錄"]
    orow = {(wso.cell(r, 7).value, str(wso.cell(r, 1).value)): [wso.cell(r, c).value for c in range(1, 12)]
            for r in range(2, wso.max_row + 1)}
    r1 = next(v for (a_, _t), v in orow.items() if a_ == "pcs_dc_current_discharge")
    check(f"舊資料：模式／單位依 action 對照、設定值與驗證結果 N/A、備註註明（{r1[3]}，{r1[4]}，{r1[5]}，{r1[8]}）",
          r1[1] == "第1次放電" and r1[3] == "直流恆流" and r1[4] == "N/A" and r1[5] == "A" and r1[8] == "N/A"
          and "舊資料未保存設定值" in r1[10])
    r2 = next(v for (a_, _t), v in orow.items() if a_ == "pcs_dc_current_charge")
    r3 = next(v for (a_, _t), v in orow.items() if a_ == "ac_on")
    check("沒有對應的充電階段 → 階段 N/A（不猜測）；非充放電控制 → 方向／模式 N/A",
          r2[1] == "N/A" and "無法可靠對應階段" in r2[10] and r3[1] == "N/A" and r3[2] == "N/A" and r3[3] == "N/A")

    print("\n== 控制紀錄：驗證結果顯示（verify_success 原值照實區分）==")
    import cert_report_xlsx as CRX
    vt = {repr(v): CRX._verify_text(v) for v in (True, False, "partial", "timeout", None, "weird", 1, "true")}
    check(f"True→成功、False→失敗、partial→部分驗證、timeout→驗證逾時、None→N/A、其他未知值→N/A：{vt}",
          vt == {"True": "成功", "False": "失敗", "'partial'": "部分驗證", "'timeout'": "驗證逾時", "None": "N/A",
                 "'weird'": "N/A", "1": "N/A", "'true'": "N/A"})
    check("備註：timeout 有說明；未知值保留原始值；True／False／None 不加備註",
          "逾時" in CRX._verify_note("timeout") and "'weird'" in CRX._verify_note("weird")
          and CRX._verify_note(True) is None and CRX._verify_note(None) is None)
    check("控制結果／整體結果的顯示邏輯不變（\"partial\" 不會被當成成功或失敗）",
          CRX._ok_text("partial") == "N/A" and CRX._ok_text(True) == "成功" and CRX._ok_text(False) == "失敗")
    root_v = os.path.join(TMP, "cert_verify")
    clk_v = Clock(t0)
    rtV = runtime(root_v, clk_v, Plant(clk_v, t0, [(0, 60, "I"), (60, 400, "D"), (400, 99999, "I")]))
    rqv = req(root_v, clk_v, "start_cert", {"cabinet_spec": "250KW"})
    rtV.tick()
    cid_v = REQ.read_result(root_v, rqv)["cert_id"]
    run_until(rtV, clk_v, 600, t0)
    base = {"cmd": "discharge", "action": "pcs_dc_current_discharge", "direction": "discharge",
            "control_mode": "dc_current", "control_mode_text": "直流恆流", "setpoint": 140.0, "unit": "A",
            "control_success": True, "success": True, "dry_run": False}
    vcases = [("T", True), ("F", False), ("P", "partial"), ("N", None), ("M", "__missing__")]
    with open(os.path.join(root_v, cid_v, "commands.jsonl"), "w", encoding="utf-8") as f:
        for k, (_tag, vv) in enumerate(vcases):
            rec_ = dict(base, t=(t0 + timedelta(seconds=55 + k)).strftime("%Y-%m-%dT%H:%M:%S"))
            if vv != "__missing__":
                rec_["verify_success"] = vv
            f.write(json.dumps(rec_, ensure_ascii=False) + "\n")
    rtV.commands = _read_jsonl(os.path.join(root_v, cid_v, "commands.jsonl"))
    rqe_v = req(root_v, clk_v, "end_cert", cert_id=cid_v)
    rtV.tick()
    wsv = load_workbook(REQ.read_result(root_v, rqe_v)["data"]["report"])["控制紀錄"]
    got = [wsv.cell(r, 9).value for r in range(2, 7)]
    check(f"報告「驗證結果」欄：True／False／partial／None／欄位不存在 → {got}",
          got == ["成功", "失敗", "部分驗證", "N/A", "N/A"]
          and "回讀僅部分確認" in str(wsv.cell(4, 11).value) and "回讀僅部分確認" not in str(wsv.cell(2, 11).value))

    print("\n== 實機 Field Test 資料（CERT_20261028_165847 複本，離線重產；不動原始資料）==")
    real = os.path.join(HERE, "..", "output", "cert_reports", "CERT_20261028_165847")
    if os.path.isdir(real):
        root_r2 = os.path.join(TMP, "cert_real")
        shutil.copytree(real, os.path.join(root_r2, "CERT_20261028_165847"),
                        ignore=shutil.ignore_patterns("reports"))
        raw_before = open(os.path.join(real, "commands.jsonl"), encoding="utf-8").read()
        rtR2 = CertRuntime(root_r2, reader=None, enable_shots=False, clock=clk_v, log=lambda *a: None)
        ok_, msg_, path_ = rtR2.generate_report("CERT_20261028_165847")
        wsr = load_workbook(path_)["控制紀錄"]
        rows_r = [[wsr.cell(r, c).value for c in range(1, 12)] for r in range(2, wsr.max_row + 1)]
        for row_ in rows_r:
            print("     ", row_[0], "|", row_[1], "|", row_[6], "| 控制", row_[7], "| 驗證", row_[8], "| 整體", row_[9])
        check("實機兩筆（放電 140 A、停止）驗證結果顯示「部分驗證」；控制結果／整體結果仍為成功；階段、設定值不變",
              ok_ and [(r_[1], r_[4], r_[5], r_[7], r_[8], r_[9]) for r_ in rows_r]
              == [("第1次放電", 140, "A", "成功", "部分驗證", "成功"), ("第1次放電", "N/A", "N/A", "成功", "部分驗證", "成功")])
        check("原始 commands.jsonl 未被修改（仍為 \"partial\"）",
              open(os.path.join(real, "commands.jsonl"), encoding="utf-8").read() == raw_before
              and raw_before.count('"verify_success": "partial"') == 2)
    else:
        print("  （略過：找不到實機認證資料夾）")

    print("\n== 設備狀態備註：待機／停機／N/A（只依 battery_power_status；不以電流 0、idle、PCS 旗標推論）==")
    from cert_stage_engine import Obs as _Obs
    ds, stt = CRX.device_state, CRX.status_text
    _t = datetime(2026, 11, 2, 9, 0, 0)

    def mk(state="休息", direction="I", stage_no=None, ok=True, **fl):
        o = _Obs(t=_t, ok=ok, values={"current_a": 0.0}, flags=dict(fl))
        o.state, o.direction, o.stage_no = state, direction, stage_no
        return o
    nm = {1: "第1次充電", 2: "第1次放電"}
    cases = [
        ("已上電＋非充放電＋無故障 → 待機", stt(mk(battery_power_status="已上電", pcs_fault_flag=False), nm), "待機"),
        ("已下電＋非充放電＋無故障 → 停機", stt(mk(battery_power_status="已下電", pcs_fault_flag=False), nm), "停機"),
        ("battery_power_status=None → N/A", stt(mk(battery_power_status=None), nm), "N/A"),
        ("battery_power_status=未知 → N/A", stt(mk(battery_power_status="未知"), nm), "N/A"),
        ("battery_power_status=切換中 → N/A", stt(mk(battery_power_status="切換中"), nm), "N/A"),
        ("讀取失敗（communication failure）→ N/A", stt(mk(ok=False, state="讀取失敗", direction=None,
                                                      battery_power_status="已上電"), nm), "N/A"),
        ("pcs_fault=True（即使已上電）→ N/A", stt(mk(battery_power_status="已上電", pcs_fault_flag=True), nm), "N/A"),
        ("本次實機異常：fault=True＋上下電未知 → N/A", stt(mk(battery_power_status="未知", pcs_fault_flag=True), nm), "N/A"),
        ("充電階段 → 第N次充電", stt(mk(state="充電", direction="C", stage_no=1, battery_power_status="已上電"), nm),
         "第1次充電"),
        ("放電階段 → 第N次放電", stt(mk(state="放電", direction="D", stage_no=2, battery_power_status="已上電"), nm),
         "第1次放電"),
        ("舊資料沒有上下電欄位＋idle → N/A（不推測待機）", stt(mk(), nm), "N/A"),
        ("非階段的突波（仍有電流方向）＋已上電 → N/A（不是待機）",
         stt(mk(state="突波", direction="D", battery_power_status="已上電"), nm), "N/A"),
        ("PCS 運行=0、待機=0 不會被當成停機（已上電 → 待機）",
         stt(mk(battery_power_status="已上電", pcs_running_flag=False, pcs_standby_flag=False), nm), "待機"),
    ]
    for lab, got_, want in cases:
        check(f"{lab}（{got_}）", got_ == want)

    root_s = os.path.join(TMP, "cert_state")
    clk_s = Clock(t0)
    # 待機（已上電）2 分 → 放電 8 分 → 待機（已上電）5 分 → 下電 3 分 → PCS 故障（上下電未知）3 分
    pl_s = Plant(clk_s, t0, [(0, 120, "I"), (120, 600, "D"), (600, 900, "I"), (900, 1080, "OFF"), (1080, 99999, "FAULT")])
    rtS2 = runtime(root_s, clk_s, pl_s)
    rq_s = req(root_s, clk_s, "start_cert", {"cabinet_spec": "250KW"})
    rtS2.tick()
    cid_s = REQ.read_result(root_s, rq_s)["cert_id"]
    run_until(rtS2, clk_s, 1260, t0)
    recs_s = _read_jsonl(os.path.join(root_s, cid_s, "observations.jsonl"))
    check("observation 保存 battery_power_status 原始值與診斷旗標（取自同一次 read_all）",
          {"battery_power_status", "pcs_running_flag", "pcs_standby_flag", "pcs_status", "pcs_fault_flag"}
          <= set(recs_s[0]["flags"]) and {r["flags"]["battery_power_status"] for r in recs_s} == {"已上電", "已下電", "未知"})
    res_s = rtS2.analysis()
    stripped = [record_to_obs(dict(r, flags={}), i + 1) for i, r in enumerate(recs_s)]
    res_plain = analyze(stripped, rtS2.p, [], [])
    check("Stage Detection 與是否保存狀態欄位無關（階段起訖、end_by、停頓完全相同）",
          [(st.no, st.dir, st.start.t, st.A.t, st.end_by, len(st.pauses)) for st in res_s.stages]
          == [(st.no, st.dir, st.start.t, st.A.t, st.end_by, len(st.pauses)) for st in res_plain.stages]
          and [o.state for o in res_s.obs] == [o.state for o in res_plain.obs])
    check("待機 ≥ 60 秒的階段結束判定不變（放電階段 end_by=idle_confirmed）",
          len(res_s.stages) == 1 and res_s.stages[0].end_by == "idle_confirmed")
    rq_es = req(root_s, clk_s, "end_cert", cert_id=cid_s)
    rtS2.tick()
    wbs = load_workbook(REQ.read_result(root_s, rq_es)["data"]["report"])
    wa_ = wbs["整體測試"]
    rmk = {c.value.strftime("%H:%M"): wa_.cell(r_, 11).value for r_, c in enumerate(wa_["A"][1:], 2)
           if hasattr(c.value, "strftime")}
    print("     整體測試備註：", {k: rmk[k] for k in ("09:01", "09:05", "09:12", "09:17", "09:20") if k in rmk})
    check("整體測試備註：待機（已上電）／第1次放電／待機／停機（已下電）／N/A（故障）",
          rmk.get("09:01") == "待機" and rmk.get("09:05") == "第1次放電" and rmk.get("09:12") == "待機"
          and rmk.get("09:17") == "停機" and rmk.get("09:20") == "N/A" and "休息" not in rmk.values())
    rws = wbs["Raw Data"]
    rh = [c.value for c in rws[1]]
    rrows = [[c.value for c in r_] for r_ in rws.iter_rows(min_row=2) if r_[3].value in ("成功", "失敗")]
    st_set = {r_[rh.index("狀態")] for r_ in rrows}
    check(f"Raw Data 狀態只出現 待機／停機／N/A／第N次放電（{sorted(st_set)}），說明保留分析狀態與上下電原始值",
          st_set == {"待機", "第1次放電", "停機", "N/A"}
          and any("分析狀態：休息" in str(r_[rh.index("說明")]) and "電池上下電：已上電" in str(r_[rh.index("說明")])
                  for r_ in rrows)
          and any("PCS 故障" in str(r_[rh.index("說明")]) for r_ in rrows))
    raw_vals = [(r_[rh.index("Cell 最大電壓")], r_[rh.index("SOC")], r_[rh.index("電流")]) for r_ in rrows]
    check("Raw Data 數值（Cell 最大電壓、SOC、電流）照原始值，未被改寫",
          raw_vals == [(r["values"]["cell_vmax_v"], r["values"]["soc"], r["values"]["current_a"]) for r in recs_s])

    print("\n== 實機舊資料（CERT_20261028_165847，沒有保存上下電）→ 非充放電 N/A ==")
    if os.path.isdir(real):
        root_o2 = os.path.join(TMP, "cert_old_state")
        shutil.copytree(real, os.path.join(root_o2, "CERT_20261028_165847"), ignore=shutil.ignore_patterns("reports"))
        rtO = CertRuntime(root_o2, reader=None, enable_shots=False, clock=clk_s, log=lambda *a: None)
        ok_o, _m, path_o = rtO.generate_report("CERT_20261028_165847")
        wbo = load_workbook(path_o)
        wo_ = wbo["整體測試"]
        orm = [(c.value.strftime("%H:%M"), wo_.cell(r_, 11).value) for r_, c in enumerate(wo_["A"][1:], 2)
               if hasattr(c.value, "strftime")]
        print("     ", orm)
        check("舊資料：放電分鐘 → 第1次放電；其他分鐘 → N/A（不回填、不推測待機）",
              ok_o and {v for _k, v in orm} == {"第1次放電", "N/A"}
              and [k for k, v in orm if v == "第1次放電"] == ["17:06", "17:07", "17:08", "17:09", "17:10"])
        ro_ = wbo["Raw Data"]
        oh = [c.value for c in ro_[1]]
        ost = [r_[oh.index("狀態")].value for r_ in ro_.iter_rows(min_row=2) if r_[3].value in ("成功", "失敗")]
        res_o = rtO._offline(os.path.join(root_o2, "CERT_20261028_165847")).analysis()
        n_dis = sum(1 for o in res_o.obs if o.state == "放電")
        check(f"舊資料 Raw Data 狀態：第1次放電 {ost.count('第1次放電')} 筆（＝分析為放電的 {n_dis} 筆）、N/A {ost.count('N/A')} 筆（共 {len(ost)}）",
              len(ost) == 84 and ost.count("第1次放電") + ost.count("N/A") == 84 and ost.count("第1次放電") == n_dis)
        check("舊資料階段判定不變（第1次放電 17:05:14 開始、idle_confirmed）",
              [(st.start.t.strftime("%H:%M:%S"), st.end_by) for st in res_o.stages] == [("17:05:14", "idle_confirmed")])
        zo = zipfile.ZipFile(path_o)
        check("舊資料：自動截圖（數據概覽／即時熱度各 1 張）、控制紀錄 2 筆仍在",
              sum(1 for n in zo.namelist() if n.startswith("xl/media/")) == 2
              and wbo["控制紀錄"].max_row == 3 and wbo["控制紀錄"].cell(2, 9).value == "部分驗證")
        zo.close()

    # 交付：模擬 Excel 副本
    if OUT_COPY:
        os.makedirs(OUT_COPY, exist_ok=True)
        for src_, name in ((reps[-1], "P3_模擬_連續記錄_含休息與重啟.xlsx"), (rr["data"]["report"], "P3_模擬_結束時仍運行_未完成.xlsx"),
                           (rep_c, "P3_模擬_控制紀錄_直流恆流與直流恆功率.xlsx")):
            shutil.copyfile(src_, os.path.join(OUT_COPY, name))
        print(f"\n模擬報告副本：{OUT_COPY}")
finally:
    CFG.CERT_MODE_ENABLED = False
    RM._report_output_root = _BAK_ROOT
    try:
        RM.release_ownership()
    except Exception:
        pass
    shutil.rmtree(TMP, ignore_errors=True)

ok = all(_results)
print(f"\n== 認證報告 P3 整合模擬 {'PASS' if ok else 'FAIL'}（{sum(_results)}/{len(_results)} 檢查通過）==")
sys.exit(0 if ok else 1)
