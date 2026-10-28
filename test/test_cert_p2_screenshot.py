# -*- coding: utf-8 -*-
"""
認證測試報告 P2 —— 截圖工作與 Excel 圖片插入（離線驗證）
========================================================
以「假的截圖函式」模擬各種結果（正常、重試、全部失敗、逾時、部分成功、資料未更新、範圍錯誤），
驗證背景執行不阻塞、紀錄 append-only、截圖類型與階段歸屬、報告選圖規則與 Excel 版面。
不連設備、不開任何網頁；實際頁面截圖另見 test_cert_p2_live.py。
執行：python test_cert_p2_screenshot.py
"""

import hashlib
import os
import re
import shutil
import struct
import sys
import tempfile
import time
import zipfile
import zlib
from datetime import datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import cert_report_config as CFG
from cert_stage_engine import Obs, analyze
from cert_report_xlsx import save_report
from cert_screenshot import (ScreenshotWorker, ShotError, caption, heat_validate, load_manifest, png_size,
                             select_stage_shot, shot_params)

_results = []


def check(label, ok):
    _results.append(bool(ok))
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}")


def make_png(w, h, rgb=(200, 220, 240)):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


SIZE = {"overview": (1088, 439), "heat": (560, 1860)}


class Fake:
    """script[target] = 依嘗試順序的結果：'ok'、('fail', 原因)、('sleep', 秒)、('stale', 落後秒)。"""

    def __init__(self, script=None, delay=0.0):
        self.script = script or {}
        self.calls = {}
        self.delay = delay

    def __call__(self, target, out_png, p, procs=None):
        n = self.calls.get(target.key, 0)
        self.calls[target.key] = n + 1
        seq = self.script.get(target.key, ["ok"])
        act = seq[min(n, len(seq) - 1)]
        time.sleep(self.delay)
        if isinstance(act, tuple) and act[0] == "fail":
            raise ShotError(act[1])
        if isinstance(act, tuple) and act[0] == "sleep":
            time.sleep(act[1])
            raise ShotError("不應到達")
        w, h = SIZE[target.key]
        with open(out_png, "wb") as f:
            f.write(make_png(w, h))
        now = datetime.now()
        lag = act[1] if isinstance(act, tuple) and act[0] == "stale" else 2
        return {"captured_at": now, "rect": [0, 0, w, h], "size": [w, h],
                "data_time": now - timedelta(seconds=lag) if target.key == "heat" else None,
                "page_clock": now - timedelta(seconds=144) if target.key == "overview" else None}


P = shot_params(SHOT_RETRY_DELAY_SEC=0.05, SHOT_ATTEMPT_HARD_TIMEOUT_SEC=1.0, SHOT_ATTEMPTS=3,
                SHOT_STALE_SEC=60)
ST1 = datetime(2026, 1, 5, 10, 0, 0)
ST2 = datetime(2026, 1, 5, 10, 25, 10)


def worker(tmp, fake, name):
    return ScreenshotWorker(os.path.join(tmp, name), params=P, capturer=fake)


def results_of(folder, job_id):
    return {r["target"]: r for r in load_manifest(folder) if r["type"] == "result" and r["job_id"] == job_id}


tmp = tempfile.mkdtemp(prefix="cert_p2_")
try:
    print("\n== 背景執行（不阻塞資料收集）==")
    w = worker(tmp, Fake(delay=0.5), "nb")
    t0 = time.monotonic()
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬：停止已確認")
    dt = time.monotonic() - t0
    ticks = 0
    while time.monotonic() - t0 < 0.4:              # 模擬資料收集迴圈持續進行
        ticks += 1
        time.sleep(0.01)
    check(f"submit() 立即回傳（{dt * 1000:.1f} ms），截圖期間主迴圈持續執行（{ticks} 次）",
          dt < 0.05 and ticks > 20)
    check("工作完成後兩個目標都有結果", w.wait_idle(10) and len(results_of(w.folder, jid)) == 2)
    rr = results_of(w.folder, jid)
    check("兩個目標並行擷取（總時間 < 兩次依序所需）", time.monotonic() - t0 < 0.95)

    print("\n== 正常截圖與紀錄內容 ==")
    w = worker(tmp, Fake(), "ok")
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬：停止已確認", trigger_time=datetime.now())
    w.wait_idle(10)
    rr = results_of(w.folder, jid)
    job = [r for r in load_manifest(w.folder) if r["type"] == "job"][0]
    check("工作紀錄含所屬階段、觸發原因、觸發時間、要求時間、類型",
          job["stage_no"] == 1 and job["stage_start"] == "2026-01-05T10:00:00"
          and job["reason"] == "模擬：停止已確認" and job["trigger_time"] and job["requested_at"]
          and job["kind"] == "stage_end")
    ov, ht = rr["overview"], rr["heat"]
    check("結果含實際擷取時間、頁面時間、檔案與尺寸",
          ov["status"] == "ok" and ov["captured_at"] and ov["page_clock"] and ht["data_time"]
          and os.path.exists(os.path.join(w.folder, ov["file"])) and ov["size"] == [1088, 439])
    check("PNG 尺寸由檔頭讀取（不需 Pillow）", png_size(os.path.join(w.folder, ht["file"])) == (560, 1860))

    print("\n== 重試、失敗、逾時、範圍錯誤、部分成功 ==")
    w = worker(tmp, Fake({"overview": [("fail", "等待逾時（30 秒）：區塊未出現"), "ok"]}), "retry")
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬")
    w.wait_idle(10)
    ov = results_of(w.folder, jid)["overview"]
    check("第一次失敗、第二次成功 → 成功，保留第一次失敗原因",
          ov["status"] == "ok" and len(ov["attempts"]) == 2 and ov["attempts"][0]["status"] == "failed"
          and "區塊未出現" in ov["attempts"][0]["reason"])

    w = worker(tmp, Fake({"overview": [("fail", "開頁失敗：net::ERR_CONNECTION_TIMED_OUT")],
                          "heat": ["ok"]}), "partial")
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬")
    w.wait_idle(10)
    rr = results_of(w.folder, jid)
    check("有限重試：全部失敗 → 嘗試 3 次後判失敗，不產生圖片",
          rr["overview"]["status"] == "failed" and len(rr["overview"]["attempts"]) == 3
          and rr["overview"]["file"] is None and "ERR_CONNECTION_TIMED_OUT" in rr["overview"]["reason"])
    check("部分成功：熱度成功、數據概覽失敗，各自記錄", rr["heat"]["status"] == "ok")
    partial = (w.folder, load_manifest(w.folder))

    w = worker(tmp, Fake({"heat": [("sleep", 3)]}), "timeout")
    t0 = time.monotonic()
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬")
    done = w.wait_idle(15)
    el = time.monotonic() - t0
    ht = results_of(w.folder, jid)["heat"]
    check(f"逾時：單次嘗試卡住 → 看門狗 1 秒判逾時，3 次後結束（{el:.1f} 秒），不卡住工作佇列",
          done and ht["status"] == "failed" and "看門狗逾時" in ht["reason"] and el < 6)

    w = worker(tmp, Fake({"heat": [("fail", "擷取範圍不符：缺少 Pack [13, 14]")]}), "range")
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬")
    w.wait_idle(10)
    check("擷取範圍錯誤 → 判失敗並記錄原因（不存不完整的圖）",
          results_of(w.folder, jid)["heat"]["status"] == "failed"
          and "擷取範圍不符" in results_of(w.folder, jid)["heat"]["reason"])

    print("\n== 資料未更新 ==")
    w = worker(tmp, Fake({"heat": [("stale", 120)]}), "stale")
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬")
    w.wait_idle(10)
    ht = results_of(w.folder, jid)["heat"]
    check("頁面更新時間落後 120 秒（> 60）→ 重試 3 次仍未更新 → 標示「資料未更新」，保留圖片",
          ht["status"] == "stale" and len(ht["attempts"]) == 3 and ht["file"]
          and "落後擷取時間 120 秒" in ht["reason"])
    stale = (w.folder, load_manifest(w.folder))
    w = worker(tmp, Fake({"heat": [("stale", 120), "ok"]}), "stale_ok")
    jid = w.submit("stage_end", 1, ST1, "第1次放電", "模擬")
    w.wait_idle(10)
    check("第一次未更新、第二次已更新 → 成功", results_of(w.folder, jid)["heat"]["status"] == "ok")

    print("\n== 截圖類型與報告選圖 ==")
    w = worker(tmp, Fake(), "kinds")
    c1 = w.submit("candidate", 1, ST1, "第1次放電", "首筆觀測待機（B）", trigger_time=datetime.now())
    c2 = w.submit("candidate", 2, ST2, "第1次充電", "首筆觀測待機（B）", trigger_time=datetime.now())
    w.wait_idle(10)
    lines_before = open(w.manifest, encoding="utf-8").read()
    w.resolve(c1, "stage_end", "待機達 60 秒，停止確認")
    w.resolve(c2, "pause", "60 秒內恢復同方向 → 停頓")
    recs = load_manifest(w.folder)
    check("判定以新紀錄附加，原紀錄不改寫（append-only）",
          open(w.manifest, encoding="utf-8").read().startswith(lines_before)
          and sum(1 for r in recs if r["type"] == "resolve") == 2)
    job, res = select_stage_shot(recs, "2026-01-05T10:00:00", "heat")
    cap = caption(job, res)
    check("候選判定為階段結束 → 採用，說明寫明「候選截圖於首筆觀測待機時擷取，停止確認後採用」",
          job["job_id"] == c1 and "候選截圖於首筆觀測待機時擷取" in cap and "實際擷取" in cap)
    job, res = select_stage_shot(recs, "2026-01-05T10:25:10", "heat")
    pz_files = [r["file"] for r in recs if r["type"] == "result" and r["job_id"] == c2]
    check("候選判定為停頓 → 不放入報告階段結束位置；圖片仍留在測試資料夾",
          job is None and res is None and len(pz_files) == 2
          and all(os.path.exists(os.path.join(w.folder, f)) for f in pz_files))

    w = worker(tmp, Fake({"overview": [("fail", "開頁失敗")]}), "delayed")
    trig = datetime.now() - timedelta(seconds=95)
    s1 = w.submit("stage_end", 1, ST1, "第1次放電", "停止已確認", trigger_time=trig)
    w.wait_idle(10)
    w.p = dict(P)
    w.capturer = Fake()
    d1 = w.submit("delayed", 1, ST1, "第1次放電", "階段結束截圖失敗後補拍", trigger_time=trig, ref_job=s1)
    w.wait_idle(10)
    recs = load_manifest(w.folder)
    job, res = select_stage_shot(recs, "2026-01-05T10:00:00", "overview")
    cap = caption(job, res)
    check("階段結束截圖失敗 → 採用延遲補拍，說明標示「非停止當下畫面」與擷取晚了多久",
          job["job_id"] == d1 and "延遲補拍：非停止當下畫面" in cap and "擷取晚 1 分 3" in cap)
    job, res = select_stage_shot(recs, "2026-01-05T10:00:00", "heat")
    check("同一工作熱度成功 → 熱度採用原階段結束截圖（各目標分開選）", job["job_id"] == s1)

    w = worker(tmp, Fake(), "prio")
    m1 = w.submit("manual", 1, ST1, "第1次放電", "操作人員手動補拍")
    w.wait_idle(10)
    s1 = w.submit("stage_end", 1, ST1, "第1次放電", "停止已確認")
    w.wait_idle(10)
    job, _ = select_stage_shot(load_manifest(w.folder), "2026-01-05T10:00:00", "overview")
    check("同時有階段結束與手動補拍 → 優先採用階段結束截圖", job["job_id"] == s1)
    check("說明文字：失敗時寫明「截圖失敗」、原因與「不使用範本」",
          all(k in caption(*select_stage_shot(partial[1], "2026-01-05T10:00:00", "overview"))
              for k in ("【截圖失敗】", "ERR_CONNECTION_TIMED_OUT", "不使用範本")))
    check("說明文字：資料未更新時標示 ⚠ 與落後秒數",
          "⚠ 頁面資料未更新" in caption(*select_stage_shot(stale[1], "2026-01-05T10:00:00", "heat")))
    bad = os.path.join(tmp, "x.png")
    open(bad, "wb").write(b"not a png")
    try:
        png_size(bad)
        check("非 PNG 檔會被拒絕", False)
    except ValueError:
        check("非 PNG 檔會被拒絕", True)

    print("\n== Excel 圖片插入 ==")
    T0 = datetime(2026, 1, 5, 10, 0, 0)
    CUR = {"D": -140.0, "C": 135.0, "I": 0.0}

    def ob(sec, k):
        return Obs(t=T0 + timedelta(seconds=sec), src=f"syn:{sec}",
                   values={"soc": 50, "current_a": CUR[k], "rack_max_c": 40, "rack_min_c": 30,
                           "voltage_v": 1200.0, "sensor_c": 25.0, "pcs_module_c": None, "power_kw": None})
    obs = ([ob(s, "D") for s in range(0, 1201, 10)] + [ob(s, "I") for s in range(1210, 1501, 10)]
           + [ob(s, "C") for s in range(1510, 2401, 10)] + [ob(s, "I") for s in range(2410, 2521, 10)]
           + [ob(s, "D") for s in range(2530, 2571, 10)] + [ob(s, "I") for s in range(2580, 2701, 10)])
    res = analyze(obs, CFG.params())
    keys = [st.start.t for st in res.stages]
    w = worker(tmp, Fake({"heat": ["ok", ("fail", "等待逾時（30 秒）：區塊數值未載入完成")]}), "xl")
    a = w.submit("stage_end", 1, keys[0], "第1次放電", "模擬：停止已確認")
    w.wait_idle(10)
    w.capturer.calls["heat"] = 1                    # 第 2 階段熱度一律失敗
    b = w.submit("stage_end", 2, keys[1], "第1次充電", "模擬：停止已確認")
    w.wait_idle(10)
    pz = w.submit("pause", 1, keys[0], "第1次放電", "模擬：停頓點")
    w.wait_idle(10)
    meta = {"conditions": ["冷氣溫度：測試", "硬體概述：測試", "測試條件：測試"], "cabinet_spec": "250KW",
            "cycles": 2, "screenshots": {"folder": w.folder, "records": load_manifest(w.folder),
                                         "note": "P2 驗證擷取，非歷史階段結束畫面"}}
    path = save_report(res, meta, os.path.join(tmp, "r.xlsx"))
    z = zipfile.ZipFile(path)
    media = sorted(n for n in z.namelist() if n.startswith("xl/media/"))
    used = {r["file"] for r in load_manifest(w.folder) if r["type"] == "result" and r["job_id"] in (a, b)
            and r["status"] == "ok"}
    want_hash = sorted(hashlib.md5(open(os.path.join(w.folder, f), "rb").read()).hexdigest() for f in used)
    got_hash = sorted(hashlib.md5(z.read(n)).hexdigest() for n in media)
    check(f"插入 3 張圖（階段1 兩張、階段2 數據概覽），內容與截圖檔完全相同；停頓點圖未插入",
          len(media) == 3 and got_hash == want_hash)

    def sheet_xml(name):
        wbx = z.read("xl/workbook.xml").decode()
        rid = re.search(rf'<sheet [^>]*name="{name}"[^>]*r:id="(\w+)"', wbx).group(1)
        rels = z.read("xl/_rels/workbook.xml.rels").decode()
        tgt = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="{rid}"', rels)
        sp = (tgt.group(1) or tgt.group(2)).lstrip("/")
        sp = sp if sp.startswith("xl/") else "xl/" + sp
        srels = z.read(sp.replace("worksheets/", "worksheets/_rels/") + ".rels").decode()
        dp = re.search(r'Target="([^"]+drawing\d+\.xml)"', srels).group(1)
        dp = "xl/drawings/" + dp.split("/")[-1]
        return z.read(dp).decode()

    def pics(dx):
        out = []
        for m in re.finditer(r"<(?:xdr:)?oneCellAnchor>(.*?)</(?:xdr:)?oneCellAnchor>", dx, re.S):
            t = m.group(1)
            if "pic" not in t:
                continue
            col = int(re.search(r"<(?:xdr:)?col>(\d+)<", t).group(1))
            row = int(re.search(r"<(?:xdr:)?row>(\d+)<", t).group(1))
            cx, cy = map(int, re.search(r'<(?:xdr:)?ext cx="(\d+)" cy="(\d+)"', t).groups())
            out.append((col, row, cx // 9525, cy // 9525))
        return out
    ovp, htp = pics(sheet_xml("數據概覽")), pics(sheet_xml("即時熱度"))
    check("圖片保持原比例與原尺寸（1088×439、560×1860，1:1 不縮放）",
          [(p_[2], p_[3]) for p_ in ovp] == [(1088, 439), (1088, 439)] and [(p_[2], p_[3]) for p_ in htp] == [(560, 1860)])
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wo, wh = wb["數據概覽"], wb["即時熱度"]
    titles = [c.row - 1 for row in wo.iter_rows(max_col=1) for c in row
              if c.value and str(c.value).endswith("結束")]
    ov_ok = all(r0 + -(-h // 20) < nxt for (_c, r0, _w, h), nxt in zip(ovp, titles[1:]))
    check("數據概覽：每張圖下緣在下一階段標題之上（不重疊）", ov_ok and len(titles) == 3)
    heat_cols = [c.column - 1 for c in wh[1] if c.value and str(c.value).endswith("結束")]
    check("即時熱度：每張圖右緣在下一階段欄位之前（以顯示縮放後較窄的 58px／欄保守計算，不重疊）",
          all(c0 + -(-w_ // 58) <= nxt for (c0, _r, w_, _h), nxt in zip(htp, heat_cols[1:])))
    texts = {ws.title: " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
             for ws in (wo, wh)}
    check("標題與圖片之間不再有多行說明（類型、擷取時間、觸發只留在 screenshots.jsonl）",
          "實際擷取" not in texts["數據概覽"] + texts["即時熱度"] and "觸發：" not in texts["數據概覽"]
          and any(r["type"] == "result" and r.get("captured_at") for r in load_manifest(w.folder)))
    check("示範身分以獨立註記列保留（放在所有圖片之後）",
          all("【註記】P2 驗證擷取，非歷史階段結束畫面" in texts[k] for k in texts)
          and max(c.row for row in wo.iter_rows() for c in row if c.value and "【註記】" in str(c.value))
          > max(r0 for _c, r0, _w, _h in ovp) + 1)
    orange = [c for ws_, cells in ((wo, [wo.cell(t + 1, 1) for t in titles]), (wh, wh[1])) for c in cells
              if c.value and str(c.value).endswith("結束")]
    check(f"階段標題：淺橘底 F4B183＋粗體（{len(orange)} 個）",
          len(orange) == 6 and all(c.fill.fgColor.rgb.endswith("F4B183") and c.font.b for c in orange))
    def fill_cols(ws_, r_, c0_):
        n_ = 0
        while ws_.cell(r_, c0_ + n_).fill.fgColor.rgb.endswith("F4B183"):
            n_ += 1
        return n_
    ov_fill = [fill_cols(wo, t + 1, 1) for t in titles]
    ht_fill = [fill_cols(wh, 1, c + 1) for c in heat_cols]
    longest = max(len(str(c.value)) for c in orange)
    check(f"標題底色只涵蓋文字＋留白：各階段統一（數據概覽 {ov_fill} 欄、即時熱度 {ht_fill} 欄），"
          f"不延伸到圖片寬度；以 58px／欄計仍容得下最長標題（{longest} 字），左側縮排留白",
          len(set(ov_fill + ht_fill)) == 1 and ov_fill[0] * 58 < 1088 / 2 and ov_fill[0] * 58 >= longest * 19
          and all(c.alignment.indent >= 1 for c in orange))
    heads = [c for c in wh[2] if c.value]
    check("即時熱度：每階段圖片上方都有粗體置中「獨立櫃Pack 即時熱度圖」，合併寬度＝圖片欄寬",
          len(heads) == 3 and all(c.value == "獨立櫃Pack 即時熱度圖" and c.font.b
                                  and c.alignment.horizontal == "center" for c in heads)
          and any(str(m) == "A2:J2" for m in wh.merged_cells.ranges) and htp[0][1] == 2)
    od = sheet_xml("數據概覽")
    arrs = re.findall(r'<xdr:absoluteAnchor[^>]*><xdr:pos x="(\d+)" y="(\d+)"/><xdr:ext cx="(\d+)" cy="(\d+)"/>'
                      r'(.*?)</xdr:absoluteAnchor>', od, re.S)
    body = arrs[0][4] if arrs else ""
    check("數據概覽：只有一支細長向下箭頭，淡藍底 8FAADC、白色粗體「時／間／序」直排",
          len(arrs) == 1 and od.count('prst="downArrow"') == 1 and 'prst="downArrow"' in body
          and 'srgbClr val="8FAADC"' in body and re.findall(r"<a:t>([^<]+)</a:t>", body) == ["時", "間", "序"]
          and body.count('b="1"') == 3 and body.count('val="FFFFFF"') == 3)
    ax, ay, acx, acy = (int(v) // 9525 for v in arrs[0][:4])

    def row_top(r0):                                # 0 起算列頂端（px）：列高以點計，不隨顯示縮放改變
        return sum(round((wo.row_dimensions[i].height or 15) * 4 / 3) for i in range(1, r0 + 1))
    img1_top = row_top(ovp[0][1])
    img2_bottom = row_top(ovp[1][1]) + ovp[1][3]
    check(f"箭頭在截圖右側留白（x {ax}～{ax + acx}px，截圖寬 {ovp[0][2]}px，與欄寬無關）、細長（寬 {acx}px），"
          f"從第 1 階段截圖頂端（{img1_top}px）到第 2 階段截圖底端（{img2_bottom}px）：{ay}～{ay + acy}px",
          ax >= ovp[0][2] + 8 and acx <= 80 and abs(ay - img1_top) <= 1 and abs(ay + acy - img2_bottom) <= 1)
    gaps = [row_top(t) - (row_top(r0) + h) for (_c, r0, _w, h), t in zip(ovp, titles[1:])]
    check(f"各階段之間的間距一致，恢復原本 2 列空白（{gaps} px），不再為箭頭加寬",
          len(set(gaps)) == 1 and 30 <= gaps[0] <= 60)
    hd = sheet_xml("即時熱度")
    check("即時熱度：已移除第 1、2 階段之間的向下箭頭（保留標題、底色、熱度圖標題與圖片）",
          'prst="downArrow"' not in hd and "absoluteAnchor" not in hd and len(htp) == 1
          and len([c for c in wh[2] if c.value == "獨立櫃Pack 即時熱度圖"]) == 3)
    cx = {}
    for n_ in sorted(x for x in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", x)):
        c_ = z.read(n_).decode()
        lo_ = float(re.search(r'<(?:c:)?min val="([^"]+)"', c_).group(1))
        ref = re.findall(r"<(?:c:)?f>([^<]+)</(?:c:)?f>", c_)
        cx[n_] = (lo_, ref)
    data_min = {}
    for ws_ in wb.worksheets:
        if ws_.title[:1].isdigit() or ws_.title == "整體測試":
            vals = [c.value for row in ws_.iter_rows(min_row=2, min_col=27, max_col=29) for c in row
                    if isinstance(c.value, (int, float))]
            if vals:
                data_min[ws_.title] = min(vals)
    margins = {}
    for n_, (lo_, ref) in cx.items():
        sheet = re.match(r"'?([^'!]+)'?!", ref[2]).group(1)
        margins[n_] = round(data_min[sheet] - lo_, 1)
    check(f"所有圖表 Y 軸下界低於資料最小值（曲線不壓在 X 軸上）：{sorted(set(margins.values()))}",
          margins and all(m > 0 for m in margins.values()))
    st_ext = [tuple(int(v) for v in m) for m in re.findall(r'<(?:xdr:)?ext cx="(\d+)" cy="(\d+)"', sheet_xml(wb.sheetnames[3]))]
    check(f"階段頁圖表尺寸同範本約 17×12.7cm（{[(round(a / 360000, 1), round(b / 360000, 1)) for a, b in st_ext]}）",
          st_ext and all(abs(a / 360000 - 17.0) < 0.05 and abs(b / 360000 - 12.7) < 0.05 for a, b in st_ext))
    check("階段2 熱度截圖失敗 → 如實顯示「截圖失敗」與原因，不放任何圖片",
          "【截圖失敗】" in texts["即時熱度"] and "區塊數值未載入完成" in texts["即時熱度"])
    check("階段3 沒有截圖紀錄 → 顯示「尚未截圖」，不使用範本",
          texts["數據概覽"].count("【尚未截圖】") == 1 and texts["即時熱度"].count("【尚未截圖】") == 1)
    sd = sheet_xml("Summary")
    check("Summary 箭頭圖形、圖表仍在（P2 不影響已確認的 Summary）",
          sd.count('prst="rightArrow"') == 1 and sd.count('prst="downArrow"') == 1
          and len(re.findall(r"<(?:xdr:)?graphicFrame[ >]", sd)) == 3)

    def emu_box(xml, prst):                         # Summary 箭頭外框（px）：A 欄 83px、其他欄 63px、列 17.4pt
        m = re.search(r"<xdr:from><xdr:col>(\d+)</xdr:col><xdr:colOff>(\d+)</xdr:colOff><xdr:row>(\d+)</xdr:row>"
                      r"<xdr:rowOff>(\d+)</xdr:rowOff></xdr:from><xdr:to><xdr:col>(\d+)</xdr:col><xdr:colOff>(\d+)"
                      r"</xdr:colOff><xdr:row>(\d+)</xdr:row><xdr:rowOff>(\d+)</xdr:rowOff></xdr:to>"
                      rf'(?:(?!</xdr:twoCellAnchor>).)*prst="{prst}"', xml, re.S)
        c0, co0, r0, ro0, c1, co1, r1, ro1 = map(int, m.groups())
        cx = lambda c, o: (0 if c == 0 else 83 + (c - 1) * 63) + o / 9525
        ry = lambda r, o: r * 17.4 * 4 / 3 + o / 9525
        return cx(c1, co1) - cx(c0, co0), ry(r1, ro1) - ry(r0, ro0), (cx(c0, co0) + cx(c1, co1)) / 2, ry(r0, ro0), ry(r1, ro1)
    hw, hh, *_ = emu_box(sd, "rightArrow")
    dw, dh, dcx, dy0, dy1 = emu_box(sd, "downArrow")
    chart_bottom = (4 + 20) * 17.4 * 4 / 3 + 214313 / 9525     # 第 1 列圖表底；下一列圖表從第 4+26 列開始
    next_top = (4 + 26) * 17.4 * 4 / 3
    check(f"Summary 向下箭頭縮小：{dw:.0f}×{dh:.0f}px（左右箭頭 {hw:.0f}×{hh:.0f}px），水平置中於右欄圖表"
          f"（中心 {dcx:.0f}px），垂直置中於上下兩圖之間（上留 {dy0 - chart_bottom:.1f}px、下留 {next_top - dy1:.1f}px）",
          abs(dw - hh) <= 20 and abs(dh - hw) <= 25 and abs(dcx - (839 + 1541) / 2) <= 3
          and dy0 > chart_bottom and dy1 < next_top and abs((dy0 - chart_bottom) - (next_top - dy1)) <= 2)
    z.close()

    print("\n== 必要狀態仍顯示；全部失敗時箭頭仍建立 ==")
    w = worker(tmp, Fake({"heat": [("stale", 120)]}), "xl_flag")
    w.submit("delayed", 1, keys[0], "第1次放電", "模擬：原截圖失敗後補拍")
    w.wait_idle(10)
    meta2 = dict(meta, screenshots={"folder": w.folder, "records": load_manifest(w.folder), "note": ""})
    wb2 = load_workbook(save_report(res, meta2, os.path.join(tmp, "r2.xlsx")))
    t2 = {ws.title: " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
          for ws in (wb2["數據概覽"], wb2["即時熱度"])}
    check("延遲補拍 → 圖片上方單行標示「非停止當下畫面」；資料未更新 → 標示「頁面資料未更新」",
          "⚠ 延遲補拍：非停止當下畫面" in t2["數據概覽"] and "頁面資料未更新" in t2["即時熱度"]
          and wb2["即時熱度"].cell(3, 1).font.color.rgb.endswith("C00000"))
    check("沒有示範註記時不產生註記列", "【註記】" not in t2["數據概覽"] + t2["即時熱度"])
    w = worker(tmp, Fake({"heat": [("fail", "開頁失敗")], "overview": [("fail", "開頁失敗")]}), "xl_allfail")
    for k_, st_ in enumerate(res.stages):
        w.submit("stage_end", st_.no, keys[k_], st_.dir_name, "模擬")
    w.wait_idle(20)
    meta3 = dict(meta, screenshots={"folder": w.folder, "records": load_manifest(w.folder), "note": ""})
    p3 = save_report(res, meta3, os.path.join(tmp, "r3.xlsx"))
    z = zipfile.ZipFile(p3)
    od3 = sheet_xml("數據概覽")
    z.close()
    wb3 = load_workbook(p3)
    check("所有截圖失敗（工作表原本沒有 drawing）→ 新建 drawing 放時間序箭頭，檔案可正常讀取",
          od3.count('prst="downArrow"') == 1 and "【截圖失敗】" in " ".join(
              str(c.value) for row in wb3["數據概覽"].iter_rows() for c in row if c.value))

    print("\n== 即時熱度區塊驗證（櫃號、Pack 集合、必要欄位與允許缺值欄位）==")

    def heat_text(cab="110", packs=range(1, 15), cell=lambda pid, i: str(40 + i % 10), sensor="-",
                  pcs=("68", "54", "50"), ext=("53", "42", "11"), others=(), n_cells=20):
        """依頁面 innerText 結構產生區塊文字（櫃名、更新時間、Sensor、PCS、Cell 極值、各 Pack 20 格＋極性）。"""
        out = []
        for c in [cab, *others]:
            out += [f"獨立櫃:{c}", "更新時間:2026-10-27 16:14:01", "Sensor", sensor, "PCS",
                    "模組溫度", pcs[0], "環境溫度", pcs[1], "櫃內溫度", pcs[2], "Cell極值",
                    "最大溫度", ext[0], "最小溫度", ext[1], "溫差", ext[2], "當前SOC", "1 %", "當前功率", "117.152 kW"]
            for pid in packs:
                out += [f"Pack {pid}"] + [cell(pid, i) for i in range(n_cells)] + ["+", "-"]
        return "\n".join(out)

    hv = heat_validate
    err, notes = hv(heat_text())
    check(f"正常：獨立櫃 110、Pack 1～14、Sensor「-」→ 通過，Sensor 只記註記：{notes}",
          err == "" and any(n.startswith("Sensor 顯示「-」") for n in notes))
    err, _ = hv(heat_text(packs=range(1, 21)))
    check(f"Pack 1～20（含 1～14）→ 判非預期 Pack，不可通過：{err}",
          err.startswith("擷取範圍不符") and "非預期 Pack 15～20" in err)
    err, _ = hv(heat_text(packs=[p for p in range(1, 15) if p != 9]))
    check(f"缺 Pack 9 → 擷取範圍不符：{err}", err.startswith("擷取範圍不符") and "缺少 Pack 9" in err)
    err, _ = hv(heat_text(packs=list(range(1, 15)) + [3]))
    check(f"Pack 3 重複 → 擷取範圍不符：{err}", "重複 Pack 3" in err and "缺少" not in err)
    err, _ = hv(heat_text(cab="109"))
    check(f"區塊屬於獨立櫃 109 → 擷取範圍不符：{err}", err.startswith("擷取範圍不符") and "獨立櫃 109" in err)
    err, _ = hv(heat_text(others=["109"]))
    check(f"區塊同時包含 110 與 109 → 擷取範圍不符：{err}", "包含其他櫃 109" in err)
    err, _ = hv(heat_text(cell=lambda pid, i: "-" if (pid, i) == (5, 3) else "45"))
    check(f"Cell 溫度缺 1 格（上限 0）→ 頁面資料無效：{err}",
          err.startswith("頁面資料無效") and "Pack 5：1 格" in err)
    err, notes = hv(heat_text(cell=lambda pid, i: "-" if (pid, i) == (5, 3) else "45"), max_missing=2)
    check("缺值上限可調整（上限 2）→ 通過並註記缺值格數", err == "" and any("缺值 1 格" in n for n in notes))
    err, _ = hv(heat_text(cell=lambda pid, i: "-", ext=("-", "-", "-"), pcs=("-", "-", "-")))
    check(f"Pack 1～14 齊全但數值全部「-」→ 不得判成功：{err}", err.startswith("頁面資料無效") and "皆為「-」" in err)
    err, _ = hv(heat_text(ext=("53", "-", "11")))
    check(f"必要欄位（最小溫度）無數值 → 頁面資料無效：{err}", err.startswith("頁面資料無效") and "最小溫度" in err)
    err, notes = hv(heat_text(pcs=("-", "-", "-")))
    check("PCS 溫度「-」屬允許缺值欄位 → 通過並各自註記", err == "" and len(notes) == 4)
    err, _ = hv(heat_text(n_cells=19))
    check(f"每 Pack 只有 19 格 → 擷取範圍不符：{err}", err.startswith("擷取範圍不符") and "20 格" in err)
    err, _ = hv(heat_text(cab="111", packs=range(1, 11)), cabinet="111", pack_ids=range(1, 11))
    check("櫃號與 Pack 數量可設定（獨立櫃 111、Pack 1～10）→ 通過", err == "")
    err, _ = hv(heat_text(), cabinet="111", pack_ids=range(1, 11))
    check("設定 111／10 Pack 時，110／14 Pack 區塊 → 不通過", err.startswith("擷取範圍不符"))

    fx = os.path.join(HERE, "fixtures")
    real = open(os.path.join(fx, "cert_heat110_20261027_data.txt"), encoding="utf-8").read()
    err, notes = hv(real)
    check(f"實際頁面文字（2026-10-27 有資料，110 區塊）→ 通過；註記：{notes}", err == "")
    nodata = open(os.path.join(fx, "cert_heat110_20261028_nodata.txt"), encoding="utf-8").read()
    err, _ = hv(nodata)
    check(f"實際頁面文字（2026-10-28 無資料，19 Pack 預設版面）→ 不通過並指出預設版面：{err}",
          err.startswith("擷取範圍不符") and "非預期 Pack 15～19" in err and "19 Pack 預設版面" in err)

    jobs = {"job_id": "j", "kind": "stage_end", "effective_kind": "stage_end", "reason": "模擬", "trigger_time": None}
    cap = caption(jobs, {"status": "ok", "target": "heat", "captured_at": "2026-10-27T16:14:03",
                         "data_time": "2026-10-27T16:14:01", "notes": ["Sensor 顯示「-」（允許缺值欄位，不影響判定）"]})
    check("說明列寫出允許缺值註記，並標明熱度頁更新時間為瀏覽器時間",
          "註：Sensor 顯示「-」" in cap and "非設備量測時間" in cap)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

ok = all(_results)
print(f"\n== 認證報告 P2 截圖（離線）{'PASS' if ok else 'FAIL'}（{sum(_results)}/{len(_results)} 檢查通過）==")
sys.exit(0 if ok else 1)
